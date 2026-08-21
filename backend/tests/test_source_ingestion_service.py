"""Task 12 — SourceIngestionService: upload/URL/process/delete orchestration
moves behind fresh compatibility hooks (SourcePipelineHooks built per call).

RED-first contracts frozen here:
- upload with a scheduler commits the queued source row BEFORE the callback
  fires and never processes inline; without a scheduler processing is inline;
- parsed source/elements commit before the best-effort chunk build;
- background element embedding overlaps extraction and 'extracted' gates on
  extraction only;
- hooks are built fresh on every call, so post-construction facade
  monkeypatches (_run_extraction) stay observed by the pipeline;
- extraction relink ordering and stale re-extraction cleanup match master;
- pipeline status/event order equals the frozen transaction_phases.json.
"""
import json
import threading
import time
import types
from pathlib import Path
from uuid import uuid4

import pytest

from app.core.config import Settings
from app.extensions import default_extension_runtime
from app.models.schemas import NotebookCreate
from app.repositories.ports import UploadedSourceFile
from app.services.source_ingestion import SourceIngestionService, SourcePipelineHooks
from app.services.sqlite_repository import SQLiteRepository, _now
from tests.model_testkit import bind_all_embedding_clients
from tests.model_testkit import bind_chat_client

ROOT = Path(__file__).resolve().parents[2]
PHASES = (
    ROOT / "backend" / "tests" / "fixtures" / "repository_contract"
    / "transaction_phases.json"
)


def _repository(settings: Settings) -> SQLiteRepository:
    return SQLiteRepository(
        settings,
        parser_provider_chain_host=default_extension_runtime().parser_chain,
    )


@pytest.fixture
def repo(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    return _repository(Settings())


@pytest.fixture
def embed_repo(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    return _repository(Settings())


class _FakeLLM:
    configured = True

    def __init__(self, payload):
        self._p = payload

    def chat_json(self, messages, response_schema_hint):
        return self._p

    def embed(self, text):
        return [0.0, 0.0]


def _element(text):
    return types.SimpleNamespace(
        element_type="paragraph", location_label="p1", text=text, metadata={}
    )


def _seed_queued_source(repo, notebook_id, file_path="/tmp/s.md"):
    sid = f"src-{uuid4().hex[:10]}"
    now = _now()
    with repo._write() as db:
        db.execute(
            "INSERT INTO sources (id,notebook_id,title,source_type,status,parse_status,"
            "file_name,file_path,file_size,file_hash,summary,doc_type,created_at,updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (sid, notebook_id, "Doc", "markdown", "queued", "queued",
             "doc.md", file_path, 0, "", "", "academic_paper", now, now))
    return sid


def test_runtime_wires_ingestion_service_and_builds_fresh_hooks(repo):
    service = repo._runtime.source_ingestion
    assert isinstance(service, SourceIngestionService)
    hooks = repo._source_pipeline_hooks()
    assert isinstance(hooks, SourcePipelineHooks)
    # Hooks are minted per call and never stored on the runtime.
    assert repo._source_pipeline_hooks() is not hooks
    assert not hasattr(repo._runtime, "source_pipeline_hooks")


def test_upload_with_scheduler_commits_queued_before_callback_and_skips_inline(
    repo, monkeypatch
):
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    inline = []
    monkeypatch.setattr(
        repo._runtime.source_ingestion,
        "process_source",
        lambda sid, hooks: inline.append(sid),
    )
    seen = []

    def scheduler(source_id):
        # A fresh connection only sees COMMITTED rows: the queued source must
        # already be durable when the scheduler callback fires.
        with repo._connect() as db:
            row = db.execute(
                "SELECT status, parse_status FROM sources WHERE id=?",
                (source_id,),
            ).fetchone()
        seen.append((source_id, row["status"], row["parse_status"]))

    out = repo.upload_sources(
        nb.id,
        [UploadedSourceFile(
            file_name="a.md", content_type="text/markdown", content=b"# T\n\nbody",
        )],
        scheduler=scheduler,
    )
    assert [s.id for s in out] == [seen[0][0]]
    assert seen[0][1:] == ("queued", "queued")
    assert inline == []
    assert out[0].parse_status == "queued"


def test_upload_without_scheduler_processes_inline(repo, monkeypatch):
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    calls = []
    monkeypatch.setattr(
        repo._runtime.source_ingestion,
        "process_source",
        lambda sid, hooks: calls.append(sid),
    )
    out = repo.upload_sources(
        nb.id,
        [UploadedSourceFile(
            file_name="a.md", content_type="text/markdown", content=b"# T\n\nbody",
        )],
    )
    assert calls == [out[0].id]


def test_parsed_source_and_elements_commit_before_chunk_build(repo, monkeypatch):
    import app.services.parser_chain_execution as parser_execution
    monkeypatch.setattr(
        parser_execution,
        "parse_builtin_source_file",
        lambda *a, **k: [_element("chunk body " * 40)],
    )
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id)
    observed = {}

    def probe(source_id):
        with repo._connect() as db:
            n = db.execute(
                "SELECT COUNT(*) c FROM source_elements WHERE source_id=?",
                (source_id,),
            ).fetchone()["c"]
            status = db.execute(
                "SELECT parse_status FROM sources WHERE id=?", (source_id,)
            ).fetchone()["parse_status"]
        observed["at_chunk_time"] = (n, status)

    monkeypatch.setattr(
        repo._runtime.source_chunking, "build_chunks_for_source", probe
    )
    repo.process_source(sid)
    assert observed["at_chunk_time"] == (1, "parsed")


def test_process_source_zeroes_chunked_at_when_writing_new_elements(
    repo, tmp_path, monkeypatch
):
    """写新代 elements 的事务内把 chunked_at 归零(旧分块完成标记随新 elements
    落库失效)。用一个真实 .md 走真实解析——**刻意不给 facade 的 parse_source_file
    打桩**(那会把测试写进冻结的 facade_surface 契约);分块步只用 runtime 探针
    (build_chunks_for_source 是运行时实例属性,打桩不进 facade 契约)在'分块时刻'
    读 chunked_at: 它在进入前非 NULL(模拟上一代已成功分块),进到分块步时已被
    写-elements 事务清成 NULL——证明归零发生在写 elements 处、早于(会重新打标的)
    分块步。"""
    md = tmp_path / "doc.md"
    md.write_text("# Heading\n\nSome body text for elements.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    # 模拟上一代已成功分块: reprocess 之前 chunked_at 是非 NULL。生产侧 chunked_at
    # 只由 replace_source_chunks 的 mark_chunked_at 随 chunk 原子提交,这里直接写这
    # 一列建 setup(无独立的 store 置方法)。
    with repo._write() as db:
        db.execute("UPDATE sources SET chunked_at=? WHERE id=?",
                   ("2020-01-01T00:00:00", sid))

    observed = {}

    def probe(source_id):
        with repo._connect() as db:
            observed["chunked_at_at_chunk_time"] = db.execute(
                "SELECT chunked_at FROM sources WHERE id=?", (source_id,)
            ).fetchone()["chunked_at"]

    monkeypatch.setattr(
        repo._runtime.source_chunking, "build_chunks_for_source", probe
    )
    repo.process_source(sid)
    # 写-elements 事务已把旧标记清零,早于(被 probe 替掉、本会重新打标的)分块步。
    assert observed["chunked_at_at_chunk_time"] is None


def test_active_lease_held_during_processing_and_released_after(
    repo, tmp_path, monkeypatch
):
    """源级活跃租约(在途误报抑制,进程内 dict):process_source 进入即 stamp 该源、
    退出(finally)即 pop。用真实 .md 走真实解析,分块步只用 **runtime 探针**
    (build_chunks_for_source 是运行时实例属性,打桩不进冻结的 facade_surface 契约)
    在'处理中'那一刻抓租约快照——该源必在;流水线跑完返回后,租约不再含它。"""
    md = tmp_path / "doc.md"
    md.write_text("# Heading\n\nSome body text for elements.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    service = repo._runtime.source_ingestion

    seen = {}

    def probe(source_id):
        # 分块时刻(进入之后、返回之前):租约本体的浅拷贝定格这一刻。
        seen["during"] = dict(service._active_sources)

    monkeypatch.setattr(
        repo._runtime.source_chunking, "build_chunks_for_source", probe
    )
    repo.process_source(sid)

    assert sid in seen["during"], "处理中租约未含该源(进入未 stamp)"
    assert sid not in service._active_sources, "退出后租约仍含该源(finally 未释放)"


def test_embedding_sources_join_active_snapshot_independently(repo):
    """codex 第5轮 P2:后台嵌入在途源(``_embedding_sources``)并入 H4/H5 用的活跃集
    (``_active_source_ids_snapshot``),但与 ``_active_sources``(process_source 生命周期 + 分块锁的
    引用计数)**独立**——process_source 可能先返回、撤自己的租约,而嵌入 worker 还在写向量;不并入
    的话在途向量被误报缺失 + 诱发重复(可能昂贵)backfill。这里直接测机制(stamp→并入快照→release
    →移出 + 与 _active_sources 隔离);process_source 在 spawn 前 stamp、_embed_bg 的 finally release
    的 wiring 见 source_ingestion。

    **变异锚点**:把 _active_source_ids_snapshot 的 ``| set(ingestion._embedding_sources)`` 去掉 →
    「snapshot 含 sid」断言红。"""
    si = repo._runtime.source_ingestion
    sid = "src-embedding-inflight"
    assert sid not in repo._runtime._active_source_ids_snapshot()

    # 模拟嵌入 worker spawn 前的 stamp(process_source 的真实做法)。
    with si._active_sources_lock:
        si._embedding_sources[sid] = si._embedding_sources.get(sid, 0) + 1
    # 并入 H4/H5 的活跃集,但不进 _active_sources(与 process_source 租约/分块锁隔离)。
    assert sid in repo._runtime._active_source_ids_snapshot(), "嵌入在途源没并进活跃集(H4/H5 会误报)"
    assert sid not in si._active_sources, "嵌入 worker 不该动 _active_sources"

    # 嵌入完成 → release → 移出快照。
    si._release_embedding_source(sid)
    assert sid not in si._embedding_sources
    assert sid not in repo._runtime._active_source_ids_snapshot()


def test_active_lease_released_on_exception_exit(repo, tmp_path, monkeypatch):
    """异常出口也释放:某阶段抛异常 → 外层 except 落 'failed' → finally 仍 pop 掉租约。

    刻意让 ``ensure_paper_metadata`` 抛异常——它是 runtime 实例方法(打桩不进 facade
    契约),在写完 elements 之后、无内层 try 兜底处调用,异常直传到外层 except。

    **变异验证锚点**:把 process_source 里 finally 的 ``self._active_sources.pop(...)``
    删掉,这条会红(异常出口后租约残留)。"""
    md = tmp_path / "doc.md"
    md.write_text("# Heading\n\nSome body text for elements.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    service = repo._runtime.source_ingestion

    def boom(*a, **k):
        raise RuntimeError("stage blew up mid-pipeline")

    monkeypatch.setattr(service, "ensure_paper_metadata", boom)

    repo.process_source(sid)  # 外层 except 吞掉异常、落 'failed' 后正常返回

    # 确证真的走了异常出口(否则这条测试会因从未触发异常而假绿)。
    assert repo.get_source(sid).parse_status == "failed"
    assert sid not in service._active_sources, "异常出口后租约残留(finally 未释放)"


def test_exhausted_guaranteed_parser_chain_marks_source_failed(
    repo, tmp_path, monkeypatch
):
    md = tmp_path / "broken.md"
    md.write_text("# unreadable", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    import app.services.parser_chain_execution as parser_execution

    monkeypatch.setattr(
        parser_execution,
        "parse_builtin_source_file",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            RuntimeError("builtin parser failed")
        ),
    )

    repo.process_source(sid)

    detail = repo.get_source(sid)
    assert detail.parse_status == "failed"
    assert detail.error_message == "parser provider chain exhausted"
    assert repo.source_elements(sid) == []


def test_active_lease_released_when_setting_parsing_status_fails(
    repo, tmp_path, monkeypatch
):
    """置 'parsing' 那句 DB 写失败也释放租约。它在 try 内首行(不在 try 外),故失败
    由外层 except 兜住落 'failed'、finally 照常 pop——若把它挪回 try 外,这条会红
    (进 try 前传出、绕过 finally,租约永久残留=体检假阴性)。

    只让 'parsing' 那次 set_source_status 抛;其余状态(尤其 except 里落 'failed')
    必须正常,否则 finally 之外的清理会被连累、测不到我们要测的那条边。"""
    md = tmp_path / "doc.md"
    md.write_text("# Heading\n\nSome body text.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    service = repo._runtime.source_ingestion

    real_set = service.set_source_status

    def fail_on_parsing(source_id, status, **kwargs):
        if status == "parsing":
            raise RuntimeError("status write hit a full disk")
        return real_set(source_id, status, **kwargs)

    monkeypatch.setattr(service, "set_source_status", fail_on_parsing)

    repo.process_source(sid)  # 外层 except 吞掉、落 'failed' 后正常返回

    assert repo.get_source(sid).parse_status == "failed"
    assert sid not in service._active_sources, (
        "置 parsing 失败后租约残留——该句仍在 try 外?"
    )


class _HardAbort(BaseException):
    """不是 Exception 子类 → 不被 process_source 的 `except Exception` 兜住,会向上
    传出。只有真正的 `finally`(而非 except 之后的正常流)能在这条出口释放租约。"""


def test_interrupted_source_lock_acquire_does_not_release_unowned_lock(
    repo, tmp_path, monkeypatch
):
    md = tmp_path / "doc.md"
    md.write_text("# Heading", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    service = repo._runtime.source_ingestion

    class InterruptedLock:
        release_calls = 0

        def acquire(self):
            raise _HardAbort("interrupted while waiting for source lock")

        def release(self):
            self.release_calls += 1
            raise AssertionError("must not release a lock that was not acquired")

    lock = InterruptedLock()
    monkeypatch.setattr(service, "_source_chunk_lock", lambda _source_id: lock)

    with pytest.raises(_HardAbort, match="interrupted while waiting"):
        repo.process_source(sid)

    assert lock.release_calls == 0
    assert sid not in service._active_sources


def test_baseexception_after_parser_materialize_cleans_uncommitted_assets(
    tmp_path, monkeypatch
):
    repo, sid = _seed_source_with_mineru_image(tmp_path, monkeypatch)
    service = repo._runtime.source_ingestion

    def hard_abort(*args, **kwargs):
        raise _HardAbort("element transaction interrupted")

    monkeypatch.setattr(service, "clear_source_extraction_state", hard_abort)

    with pytest.raises(_HardAbort, match="element transaction interrupted"):
        repo.process_source(sid)

    assert repo.source_asset_ids(sid) == []


def test_concurrent_same_source_reparse_serializes_asset_generations(
    tmp_path, monkeypatch
):
    repo, sid = _seed_source_with_mineru_image(tmp_path, monkeypatch)
    entered = threading.Event()
    release = threading.Event()
    calls = []
    calls_lock = threading.Lock()
    png = b"\x89PNG\r\n\x1a\n" + b"1" * 32

    def parse_once(path, name):
        with calls_lock:
            calls.append(name)
            first = len(calls) == 1
        if first:
            entered.set()
            assert release.wait(timeout=5)
        return (
            [
                {
                    "type": "image",
                    "img_path": "fig.png",
                    "image_caption": ["Figure."],
                    "page_idx": 0,
                }
            ],
            {"fig.png": png},
        )

    monkeypatch.setattr(repo.mineru_client, "parse_with_images", parse_once)
    errors = []

    def run():
        try:
            repo.process_source(sid)
        except BaseException as exc:  # test thread must report every failure
            errors.append(exc)

    first = threading.Thread(target=run)
    second = threading.Thread(target=run)
    first.start()
    assert entered.wait(timeout=5)
    second.start()
    time.sleep(0.1)
    assert len(calls) == 1
    release.set()
    first.join(timeout=10)
    second.join(timeout=10)

    assert not first.is_alive()
    assert not second.is_alive()
    assert errors == []
    assert len(calls) == 2
    assert len(repo.source_asset_ids(sid)) == 1


def test_active_lease_released_on_propagating_baseexception(
    repo, tmp_path, monkeypatch
):
    """「移动」变异防线。把 pop 从 finally 挪到 except 之后的正常流,对**被捕获**的
    异常照样会跑(except 不重抛)——那种搬移下 test_active_lease_released_on_exception_exit
    仍会绿,漏掉这条搬移。这条用一个不是 Exception 子类的 BaseException:它绕过外层
    `except Exception` 向上传出,只有 finally 能在传出出口释放租约。

    双重变异锚点:(a) 删掉 finally 的 pop → 红;(b) 把 pop 挪出 finally 到 except
    之后 → 传出出口跳过它 → 红。"""
    md = tmp_path / "doc.md"
    md.write_text("# Heading\n\nSome body text for elements.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    service = repo._runtime.source_ingestion

    def hard_abort(*a, **k):
        raise _HardAbort("propagating past except Exception")

    # ensure_paper_metadata 在 embed 线程启动之前调用,传出不会遗留后台线程。
    monkeypatch.setattr(service, "ensure_paper_metadata", hard_abort)

    with pytest.raises(_HardAbort):
        repo.process_source(sid)

    assert sid not in service._active_sources, "传出出口后租约残留(pop 不在 finally)"


def test_active_lease_is_refcounted_across_overlapping_invocations(
    repo, tmp_path, monkeypatch
):
    """引用计数:同一源被并发处理时(上传后台 job 未完时 owner 又点 parse),先完成的
    invocation 不能撤掉仍在跑者的租约。用 probe 在'处理中'那一刻模拟第二个并发
    invocation 也进入(给计数再加一,1→2);第一个 process_source 返回、finally 减一后
    (2→1),租约必须仍含该源——否则仍在跑的第二个会在途缺 elements/chunks 被体检误报
    损坏。

    **变异锚点**:把 finally 的「减一、归零才 pop」改回无条件 ``pop(source_id, None)``,
    这条会红(先完成者把仍在跑者的租约一起撤了,计数语义退化成布尔在场)。"""
    md = tmp_path / "doc.md"
    md.write_text("# Heading\n\nSome body text for elements.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    service = repo._runtime.source_ingestion

    def probe(source_id):
        # 模拟另一个并发 invocation 也进入了(第二次 stamp,计数 1→2)。
        with service._active_sources_lock:
            service._active_sources[source_id] = (
                service._active_sources.get(source_id, 0) + 1
            )

    monkeypatch.setattr(
        repo._runtime.source_chunking, "build_chunks_for_source", probe
    )
    repo.process_source(sid)

    # 第一个 invocation 的 finally 把计数 2→1,租约仍在(模拟的第二个还没 finally)。
    assert sid in service._active_sources, "先完成的 invocation 撤掉了仍在跑者的租约"
    assert service._active_sources[sid] == 1
    # 清理模拟的第二个 invocation 的租约,免得污染后续用例。
    with service._active_sources_lock:
        service._active_sources.pop(sid, None)


def test_concurrent_same_source_reparse_serializes_element_swap_and_chunk_build(
    repo, tmp_path, monkeypatch
):
    """per-source 分块锁:同源两次并发 process_source 必须在「换 elements → 建 chunks
    + 置 marker」**整段**串行、不得交错——否则一次 invocation 读到 A 代 elements、另一次
    换成 B 代、第一次再把 A 代 chunks 连同 marker 写回,留下「B 代 elements + A 代 chunks
    + marker 已置」的假完成(codex 第 2 轮 P2)。

    检测的是**跨 invocation 的临界区所有权交错**,而非仅仅 build_chunks 重叠——后者
    只锁 build_chunks(不含 replace_elements)也能满足,却仍会被 replace_elements 的
    交错破坏(移动变异)。故这里在 replace_elements 处认领所有权、在 build_chunks 结束
    处释放:一旦某 invocation 认领了区间,另一 invocation 又跑到 replace_elements 就是
    交错。停顿只放大「窄锁/无锁」变异的暴露窗口;WITH 正确锁时另一线程阻塞在锁上、
    根本到不了 replace_elements,与时序无关,故 CI 不 flaky。

    **双变异锚点**:(a) 去掉 `with self._source_chunk_lock(source_id):` 包裹 → 红;
    (b) 把包裹缩小到只裹 build_chunks(不含 replace_elements 的 write 块)→ 仍红
    (replace_elements 交错被认领检查抓到)。"""
    import contextvars

    md = tmp_path / "doc.md"
    md.write_text("# Heading\n\nSome body text for elements.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    store = repo._runtime.source_store

    owner = {"tid": None}
    interleaved = {"seen": False}
    probe_lock = threading.Lock()
    real_replace = store.replace_elements
    real_build = repo._runtime.source_chunking.build_chunks_for_source

    def probe_replace(db, source_id, elements, *, created_at):
        # 认领临界区:若此刻另一 invocation 已认领且尚未释放,即为交错。
        with probe_lock:
            if owner["tid"] not in (None, threading.get_ident()):
                interleaved["seen"] = True
            owner["tid"] = threading.get_ident()
        return real_replace(db, source_id, elements, created_at=created_at)

    def probe_build(source_id):
        time.sleep(0.05)  # 放大 replace→build 区间,让窄锁/无锁变异可靠交错
        real_build(source_id)
        with probe_lock:  # 释放临界区所有权(区间结束)
            owner["tid"] = None

    monkeypatch.setattr(store, "replace_elements", probe_replace)
    monkeypatch.setattr(
        repo._runtime.source_chunking, "build_chunks_for_source", probe_build
    )

    def run():
        repo.process_source(sid)

    # copy_context:裸线程要带上当前上下文(current-user 等 ContextVar),否则管线里
    # 读上下文的分支会拿到错的默认值(见 user-system-state 教训)。
    threads = [
        threading.Thread(target=contextvars.copy_context().run, args=(run,))
        for _ in range(2)
    ]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert not interleaved["seen"], (
        "并发同源 reparse 的「换 elements→建 chunks」区间交错了(串行锁未覆盖 replace)"
    )
    # 收尾一致性:两次成功 reparse 串行后 marker 应已置(且是后完成者的世代)。
    with repo._connect() as db:
        chunked_at = db.execute(
            "SELECT chunked_at FROM sources WHERE id=?", (sid,)
        ).fetchone()["chunked_at"]
    assert chunked_at is not None, "两次成功 reparse 串行后 marker 应已置"
    # 锁在 refcount 归零后被清除(有界化,不泄漏)。
    assert sid not in repo._runtime.source_ingestion._source_chunk_locks


def test_backfill_lock_guard_registers_lease_so_lock_survives_concurrent_release(repo):
    """codex 第2轮 P1:hold_source_chunk_lock(体检 backfill 用)持锁期间**登记活跃租约**,使并发
    process_source 的 finally(_release_source_lease)不会在本方仍持锁时把锁 pop 掉——否则后续
    reparse 会另建一把**新锁**、与本方互斥失效,给复用的 element/chunk id 挂上永久陈旧向量;且
    backfill-only 的源没有租约触发 pop→锁泄漏。本方作为最后持有者退出时连锁带租约一并清除。

    **变异锚点**:把 hold_source_chunk_lock 里登记租约的那句 stamp 去掉(或改回直接
    `with self._source_chunk_lock(...)`)→ 持锁期租约为 0,下面模拟的并发释放会把锁 pop 掉 →
    `is held`(锁仍在且是同一把)断言变红。"""
    si = repo._runtime.source_ingestion
    sid = "src-backfill-hold"
    assert sid not in si._active_sources and sid not in si._source_chunk_locks
    with si.hold_source_chunk_lock(sid):
        held = si._source_chunk_locks.get(sid)
        assert held is not None and held.locked()       # 锁在注册表且被本方持有
        assert si._active_sources.get(sid) == 1          # 持锁期登记了租约
        # 模拟一个并发 process_source:先 stamp(1→2),干完 finally 减一(2→1)。
        with si._active_sources_lock:
            si._active_sources[sid] = si._active_sources.get(sid, 0) + 1
        si._release_source_lease(sid)
        assert si._source_chunk_locks.get(sid) is held   # 本方租约还在 → 锁没被 pop(同一把)
        assert si._source_chunk_locks[sid].locked()      # 仍被本方持有
    # 退出:本方是最后持有者 → 锁与租约都清除(有界化,不泄漏)。
    assert sid not in si._active_sources
    assert sid not in si._source_chunk_locks


class _ElementBlockingEmbedder:
    """Blocks ONLY element-embedding worker threads (named 'emb-el*'), so KG
    object embedding inside extraction proceeds while element embedding is
    held — proving 'extracted' waits for extraction only."""

    def __init__(self, dim=8):
        self.dim = dim
        self.entered = threading.Event()
        self.release = threading.Event()

    def embed_texts(self, texts):
        if threading.current_thread().name.startswith("emb-el"):
            self.entered.set()
            self.release.wait(15)
        return [[0.1] * self.dim for _ in texts]

    def embed_query(self, text):
        return [0.0] * self.dim

    def _ensure(self):
        pass


def test_background_embedding_overlaps_extraction_and_extracted_gates_on_extraction(
    embed_repo, tmp_path
):
    repo = embed_repo
    md = tmp_path / "doc.md"
    md.write_text("# Doc\n\nEngram is a memory architecture.\n", encoding="utf-8")
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id, file_path=str(md))
    repo.settings.kg_auto_extract = True
    bind_chat_client(repo, "kg_extract", _FakeLLM(json.dumps({
        "nodes": [{"local_id": "a", "type": "Concept", "name": "Engram",
                   "evidence": "Engram is a memory architecture"}],
        "edges": []})))
    emb = _ElementBlockingEmbedder()
    bind_all_embedding_clients(repo, emb)

    done = threading.Event()

    def run():
        try:
            repo.process_source(sid)
        finally:
            done.set()

    threading.Thread(target=run, daemon=True).start()
    assert emb.entered.wait(10), "background element-embedding never started"

    deadline = time.time() + 10
    reached = False
    while time.time() < deadline:
        if repo.get_source(sid).parse_status == "extracted":
            reached = True
            break
        time.sleep(0.05)
    assert not emb.release.is_set(), "precondition: embedding still blocked"
    assert reached, "'extracted' must not wait for the background element embed"

    emb.release.set()
    assert done.wait(15), "process_source did not finish after releasing embedder"
    with repo._connect() as db:
        (n,) = db.execute(
            "SELECT COUNT(*) FROM element_embeddings WHERE source_id=?", (sid,)
        ).fetchone()
    assert n >= 1, "element embeddings must persist once the pipeline completes"


def test_fresh_hooks_preserve_post_construction_component_monkeypatch(
    repo, monkeypatch
):
    import app.services.parser_chain_execution as parser_execution
    monkeypatch.setattr(
        parser_execution,
        "parse_builtin_source_file",
        lambda *a, **k: [_element("body text")],
    )
    repo.settings.kg_auto_extract = True
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id)
    calls = []
    monkeypatch.setattr(repo._runtime.source_ingestion, "run_extraction", lambda sid: calls.append(sid))
    repo.process_source(sid)
    assert calls == [sid], "hooks must re-resolve the component seam on every call"
    assert repo.get_source(sid).parse_status == "extracted"


def test_extraction_relink_ordering_and_stale_source_cleanup_match_master(repo):
    """Relink edges are proposed BEFORE store_kg (same review_status/source_id
    remap as LLM edges) and a re-extraction replaces — never duplicates —
    the source's prior objects."""
    bind_chat_client(repo, "kg_extract", _FakeLLM(json.dumps({
        "nodes": [
            {"local_id": "a", "type": "Concept", "name": "Engram",
             "evidence": "Engram is a memory architecture"},
            {"local_id": "b", "type": "Claim", "name": "Engram improves perplexity",
             "evidence": "Engram improves perplexity"},
        ],
        "edges": []})))
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = f"src-{uuid4().hex[:10]}"
    now = _now()
    with repo._write() as db:
        db.execute(
            "INSERT INTO sources (id,notebook_id,title,source_type,status,parse_status,"
            "file_name,file_path,file_size,file_hash,summary,doc_type,created_at,updated_at) "
            "VALUES (?,?,?, 'markdown','extracted','parsed', 'doc.md','',0,'','',"
            "'academic_paper',?,?)",
            (sid, nb.id, "Doc", now, now))
        db.execute(
            "INSERT INTO source_elements (id,source_id,element_type,location_label,"
            "text,metadata,created_at) VALUES (?,?,?,?,?,?,?)",
            (f"el-{sid}-0001", sid, "paragraph", "p1",
             "Engram is a memory architecture. Engram improves perplexity.",
             "{}", now))
    repo._run_extraction(sid)
    rels = repo.relations_for_notebook(nb.id)
    assert len(rels) == 1, f"deterministic relink must reconnect degree-0 nodes: {rels}"
    with repo._connect() as db:
        obj_ids = {r["id"] for r in db.execute(
            "SELECT id FROM knowledge_objects WHERE notebook_id=?", (nb.id,)
        ).fetchall()}
        (n1,) = db.execute(
            "SELECT COUNT(*) FROM knowledge_objects WHERE source_id=?", (sid,)
        ).fetchone()
    assert rels[0]["source_object_id"] in obj_ids
    assert rels[0]["target_object_id"] in obj_ids

    repo._run_extraction(sid)  # stale-source cleanup: replaced, not duplicated
    with repo._connect() as db:
        (n2,) = db.execute(
            "SELECT COUNT(*) FROM knowledge_objects WHERE source_id=?", (sid,)
        ).fetchone()
    assert n2 == n1
    assert len(repo.relations_for_notebook(nb.id)) == 1


def test_pipeline_status_and_event_order_equals_transaction_phases(repo, monkeypatch):
    frozen = json.loads(PHASES.read_text(encoding="utf-8"))["process_source"]
    assert frozen["sequence"] == [
        "set parsing",
        "parse outside transaction",
        "replace elements and source-derived state in one write",
        "set parsed",
        "best-effort chunk build",
        "background embedding",
        "foreground extraction",
        "set extracted",
        "join embedding thread",
        "enqueue existing-index fold",
    ]

    import app.services.parser_chain_execution as parser_execution
    monkeypatch.setattr(
        parser_execution,
        "parse_builtin_source_file",
        lambda *a, **k: [_element("event body")],
    )
    repo.settings.kg_auto_extract = True
    monkeypatch.setattr(
        repo._runtime.source_ingestion, "run_extraction", lambda sid: None
    )
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = _seed_queued_source(repo, nb.id)

    events = []
    original_emit = repo.event_log.emit
    monkeypatch.setattr(
        repo.event_log,
        "emit",
        lambda payload, **kw: (events.append(dict(payload)), original_emit(payload, **kw))[0],
    )
    repo.process_source(sid)

    labels = []
    for event in events:
        if event.get("kind") == "status":
            labels.append(f"status:{event['status']}")
        elif event.get("kind") == "pipeline":
            labels.append(f"pipeline:{event['stage']}:{event['status']}")

    def index(label):
        assert label in labels, (label, labels)
        return labels.index(label)

    # Order: parsing → parse → parsed → extracting → embed start (background) →
    # extract → extracted → embed joined → pipeline done last. The parsed→extracting
    # status flip is emitted before the background embed thread is spawned, so
    # status:extracting precedes embed:start.
    assert index("status:parsing") < index("pipeline:parse:start")
    assert index("pipeline:parse:start") < index("pipeline:parse:done")
    assert index("pipeline:parse:done") < index("status:parsed")
    assert index("status:parsed") < index("status:extracting")
    assert index("status:extracting") < index("pipeline:embed:start")
    assert index("pipeline:embed:start") < index("pipeline:extract:start")
    assert index("pipeline:extract:start") < index("pipeline:extract:done")
    assert index("pipeline:extract:done") < index("status:extracted")
    assert index("pipeline:embed:done") < index("pipeline:pipeline:done")
    assert labels[-1] == "pipeline:pipeline:done"


def test_make_persist_image_enforces_guardrails(monkeypatch, tmp_path):
    # facade 工厂：超尺寸/超张数被挡；关配置时整体返回 None。
    from app.core.config import Settings
    from app.services.knowhow.assets import AssetService

    class _Repo:
        storage_dir = tmp_path
        def __init__(self): self.saved = []
        def insert_notebook_asset(self, nb, fn, mime, size, by, source_id=None):
            aid = f"a{len(self.saved)}"; self.saved.append(aid); return aid
        def get_notebook_asset(self, aid): return {"id": aid, "notebook_id": "nb", "mime": "image/png"}

    # 直接测工厂逻辑：构造与 facade._make_persist_image 等价的闭包工厂
    from app.services.source_image_persist import make_persist_image_factory  # Task 引入的纯函数
    repo = _Repo()
    settings = Settings(MINERU_RETURN_IMAGES="1", MINERU_MAX_IMAGE_BYTES="10",
                        MINERU_MAX_IMAGES_PER_SOURCE="1")
    factory = make_persist_image_factory(settings, lambda: AssetService(repo))
    persist = factory("nb", "src-1", "u")
    assert persist(b"x" * 5, "a.png") == "a0"      # 通过
    assert persist(b"x" * 50, "b.png") is None      # 超尺寸被挡
    assert persist(b"x" * 5, "c.png") is None        # 超张数(上限1)被挡

    off = Settings(MINERU_RETURN_IMAGES="0")
    assert make_persist_image_factory(off, lambda: AssetService(repo))("nb", "s", "u") is None


def test_process_source_persists_mineru_local_image_to_source_assets(tmp_path, monkeypatch):
    """Task 8 end-to-end: an image block returned by (fake) local MinerU,
    flowing through the REAL facade (process_source -> make_persist_image ->
    parse_file -> parse_pdf -> mineru_content_list_to_elements ->
    AssetService.save_source_image), actually lands as a notebook_assets row
    tagged with this source's id AND a real file on disk with matching
    bytes — the behavior this whole task exists to wire up, beyond the
    guardrail-only unit test above."""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.setenv("MINERU_MODE", "http")
    monkeypatch.setenv("MINERU_API_URL", "http://localhost:8888")
    repo = _repository(Settings())

    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = f"src-{uuid4().hex[:10]}"
    now = _now()
    with repo._write() as db:
        db.execute(
            "INSERT INTO sources (id,notebook_id,title,source_type,status,parse_status,"
            "file_name,file_path,file_size,file_hash,summary,doc_type,created_at,updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (sid, nb.id, "Doc", "pdf", "queued", "queued",
             "doc.pdf", "/tmp/doc.pdf", 0, "", "", "academic_paper", now, now))

    png_bytes = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    monkeypatch.setattr(
        repo.mineru_client, "parse_with_images",
        lambda path, name: (
            [{"type": "image", "img_path": "fig1.png",
              "image_caption": ["Figure 1."], "page_idx": 0}],
            {"fig1.png": png_bytes},
        ),
    )
    repo.process_source(sid)

    detail = repo.get_source(sid)
    assert detail.parse_status == "extracted"
    asset_ids = repo.source_asset_ids(sid)
    assert len(asset_ids) == 1
    asset = repo.get_notebook_asset(asset_ids[0])
    assert asset["notebook_id"] == nb.id

    from app.services.knowhow.assets import AssetService
    on_disk = AssetService(repo).path_for(asset)
    assert on_disk.is_file()
    assert on_disk.read_bytes() == png_bytes

    elements = repo.source_elements(sid)
    image_el = next(e for e in elements if e.element_type == "image")
    assert image_el.metadata.get("asset_id") == asset_ids[0]


def _seed_source_with_mineru_image(tmp_path, monkeypatch):
    """Shared setup (Task 8's e2e style): a real facade + a queued pdf source
    whose (fake) local MinerU parse yields one image block, so process_source
    persists exactly one notebook_assets row + on-disk file. Used by both
    Task 9 tests below (re-parse cleanup / delete cleanup) to get to a
    just-parsed source with a real image asset in one shared call."""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.setenv("MINERU_MODE", "http")
    monkeypatch.setenv("MINERU_API_URL", "http://localhost:8888")
    repo = _repository(Settings())

    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = f"src-{uuid4().hex[:10]}"
    now = _now()
    with repo._write() as db:
        db.execute(
            "INSERT INTO sources (id,notebook_id,title,source_type,status,parse_status,"
            "file_name,file_path,file_size,file_hash,summary,doc_type,created_at,updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (sid, nb.id, "Doc", "pdf", "queued", "queued",
             "doc.pdf", "/tmp/doc.pdf", 0, "", "", "academic_paper", now, now))

    png_bytes = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    monkeypatch.setattr(
        repo.mineru_client, "parse_with_images",
        lambda path, name: (
            [{"type": "image", "img_path": "fig1.png",
              "image_caption": ["Figure 1."], "page_idx": 0}],
            {"fig1.png": png_bytes},
        ),
    )
    return repo, sid


def test_process_source_reparse_clears_stale_source_images(tmp_path, monkeypatch):
    """Task 9: re-parsing the SAME source (parse_status already 'extracted')
    must not accumulate MinerU image assets across runs. process_source now
    calls delete_source_images(source_id) before the new parse persists its
    own images, so a second process_source call leaves exactly ONE asset row
    (the fresh one) — never two (stale + fresh) — and the stale row/file are
    actually gone, not merely uncounted."""
    repo, sid = _seed_source_with_mineru_image(tmp_path, monkeypatch)

    repo.process_source(sid)
    first_ids = repo.source_asset_ids(sid)
    assert len(first_ids) == 1

    from app.services.knowhow.assets import AssetService
    asset_svc = AssetService(repo)
    first_asset = repo.get_notebook_asset(first_ids[0])
    first_path = asset_svc.path_for(first_asset)
    assert first_path.is_file()

    # Re-parse the same source (e.g. user clicks "重新解析"). Without the
    # Task 9 cascade-delete this would append a second row/file and double
    # the count.
    repo.process_source(sid)
    second_ids = repo.source_asset_ids(sid)
    assert len(second_ids) == 1
    assert second_ids != first_ids  # old row replaced, not appended to

    assert repo.get_notebook_asset(first_ids[0]) is None  # stale row gone
    assert not first_path.is_file()  # stale file gone, not just orphaned


def test_delete_source_cleans_images(tmp_path, monkeypatch):
    """Task 9: deleting a source must cascade-delete its MinerU image assets
    (notebook_assets rows + on-disk files) alongside its elements/file,
    otherwise every deleted source with embedded images leaves orphaned
    asset rows and files behind forever."""
    repo, sid = _seed_source_with_mineru_image(tmp_path, monkeypatch)
    repo.process_source(sid)
    asset_ids = repo.source_asset_ids(sid)
    assert len(asset_ids) == 1

    from app.services.knowhow.assets import AssetService
    asset = repo.get_notebook_asset(asset_ids[0])
    on_disk = AssetService(repo).path_for(asset)
    assert on_disk.is_file()

    repo.delete_source(sid)

    assert repo.source_asset_ids(sid) == []
    assert repo.get_notebook_asset(asset_ids[0]) is None
    assert not on_disk.is_file()


def _seed_queued_pdf(repo, tmp_path):
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = f"src-{uuid4().hex[:10]}"
    now = _now()
    with repo._write() as db:
        db.execute(
            "INSERT INTO sources (id,notebook_id,title,source_type,status,parse_status,"
            "file_name,file_path,file_size,file_hash,summary,doc_type,created_at,updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (sid, nb.id, "Doc", "pdf", "queued", "queued",
             "doc.pdf", "/tmp/doc.pdf", 0, "", "", "academic_paper", now, now))
    return nb, sid


def test_upload_file_uses_cloud_when_only_cloud_configured(tmp_path, monkeypatch):
    """本地 MinerU 未配置 + 云端已配 → 上传文件走云端 parse_file_with_images，
    图片作为 notebook_assets 落地(与本地路径同一持久化闭包)。"""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.delenv("MINERU_MODE", raising=False)          # 本地 off → not configured
    monkeypatch.delenv("MINERU_API_URL", raising=False)
    monkeypatch.setenv("MINERU_API_TOKEN", "tok")             # 云端 configured
    repo = _repository(Settings())
    nb, sid = _seed_queued_pdf(repo, tmp_path)

    png = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    monkeypatch.setattr(
        repo.mineru_cloud_client, "parse_file_with_images",
        lambda path, data_id="": (
            [{"type": "image", "img_path": "fig1.png",
              "image_caption": ["Figure 1."], "page_idx": 0}],
            {"fig1.png": png},
        ),
    )
    repo.process_source(sid)

    assert repo.get_source(sid).parse_status == "extracted"
    asset_ids = repo.source_asset_ids(sid)
    assert len(asset_ids) == 1
    elements = repo.source_elements(sid)
    assert any(e.element_type == "image" for e in elements)


def test_upload_file_cloud_error_falls_back_to_local_python(tmp_path, monkeypatch):
    """云端上传抛错 → 本地 Python 解析；空白页仍保留降级质量警告。"""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.delenv("MINERU_MODE", raising=False)
    monkeypatch.delenv("MINERU_API_URL", raising=False)
    monkeypatch.setenv("MINERU_API_TOKEN", "tok")
    repo = _repository(Settings())
    nb, sid = _seed_queued_pdf(repo, tmp_path)

    def _boom(path, data_id=""):
        raise RuntimeError("云端 500")
    monkeypatch.setattr(repo.mineru_cloud_client, "parse_file_with_images", _boom)
    # pypdf 回落读取真实文件：写一个最小 PDF 到 source.file_path
    import pypdf, io as _io
    writer = pypdf.PdfWriter(); writer.add_blank_page(width=200, height=200)
    real_pdf = tmp_path / "real.pdf"
    with real_pdf.open("wb") as fh: writer.write(fh)
    with repo._write() as db:
        db.execute("UPDATE sources SET file_path=? WHERE id=?", (str(real_pdf), sid))

    repo.process_source(sid)   # 不抛：云端错 → 本地 Python PDF 兜底
    detail = repo.get_source(sid)
    assert detail.parse_status == "extracted"
    assert detail.parse_quality_warning is True
    assert detail.error_message.startswith("[pdf-python-fallback]")
    assert "No extractable text" in detail.error_message
    assert repo.source_asset_ids(sid) == []   # 空白页无图片资产


# -- 降级质量警告泛化到全部 MinerU 可解析后缀 --------------------------------
#
# 判据此前钉死 `.pdf`，于是 docx/pptx/xlsx 走完 MinerU 优先分支再回落本地库时
# 是**静默**的：用户拿到的是低保真结果，界面上却没有任何提示。


def _seed_queued_office(repo, file_path, file_name, source_type):
    nb = repo.create_notebook(NotebookCreate(name="nb"))
    sid = f"src-{uuid4().hex[:10]}"
    now = _now()
    with repo._write() as db:
        db.execute(
            "INSERT INTO sources (id,notebook_id,title,source_type,status,parse_status,"
            "file_name,file_path,file_size,file_hash,summary,doc_type,created_at,updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (sid, nb.id, "Doc", source_type, "queued", "queued",
             file_name, str(file_path), 0, "", "", "academic_paper", now, now))
    return nb, sid


def _write_docx(path):
    from docx import Document

    document = Document()
    document.add_paragraph("Hello from docx fallback.")
    document.save(str(path))
    return path


def test_docx_local_python_fallback_after_mineru_error_warns(tmp_path, monkeypatch):
    """MinerU 已配置但产出本机兜底链（mammoth → python-docx）→ 稳定前缀 + 质量警告。"""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.setenv("MINERU_MODE", "http")                 # 本地 MinerU configured
    monkeypatch.setenv("MINERU_API_URL", "http://mineru.internal")
    monkeypatch.delenv("MINERU_API_TOKEN", raising=False)
    repo = _repository(Settings())
    docx_path = _write_docx(tmp_path / "real.docx")
    nb, sid = _seed_queued_office(repo, docx_path, "real.docx", "docx")

    def _boom(file_path, file_name):
        raise RuntimeError("mineru 502")
    monkeypatch.setattr(repo.mineru_client, "parse_with_images", _boom)

    repo.process_source(sid)
    detail = repo.get_source(sid)
    assert detail.parse_status == "extracted"          # 降级仍是成功摄取
    assert detail.parse_quality_warning is True
    assert detail.error_message.startswith("[pdf-python-fallback]")
    assert "PDF" not in detail.error_message.removeprefix("[pdf-python-fallback]")
    assert any("Hello from docx fallback." in e.text for e in repo.source_elements(sid))


def test_docx_fallback_without_mineru_configured_does_not_warn(tmp_path, monkeypatch):
    """MinerU 根本没配置时的正常兜底不是降级，不许打警告（否则每个纯本地部署
    的每份 docx 都常亮一个假异常）。"""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.delenv("MINERU_MODE", raising=False)
    monkeypatch.delenv("MINERU_API_URL", raising=False)
    monkeypatch.delenv("MINERU_API_TOKEN", raising=False)
    repo = _repository(Settings())
    docx_path = _write_docx(tmp_path / "real.docx")
    nb, sid = _seed_queued_office(repo, docx_path, "real.docx", "docx")

    repo.process_source(sid)
    detail = repo.get_source(sid)
    assert detail.parse_status == "extracted"
    assert detail.parse_quality_warning is False
    assert not (detail.error_message or "").startswith("[pdf-python-fallback]")


def _write_xlsx(path):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["Header A", "Header B"])
    sheet.append(["Hello from xlsx fallback.", "42"])
    book.save(str(path))
    return path


def test_xlsx_fallback_to_openpyxl_does_not_warn(tmp_path, monkeypatch):
    """工作簿刻意不在警告集里：openpyxl 兜底对单元格值是全保真的（拿的是逐格原
    值，不是版面重建），打「质量降级」是噪音；而部署的 MinerU 若不支持工作簿，
    这条警告会在每份表格上永远点亮、重新解析也消不掉。"""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.setenv("MINERU_MODE", "http")                 # 本地 MinerU configured
    monkeypatch.setenv("MINERU_API_URL", "http://mineru.internal")
    monkeypatch.delenv("MINERU_API_TOKEN", raising=False)
    repo = _repository(Settings())
    book = _write_xlsx(tmp_path / "real.xlsx")
    nb, sid = _seed_queued_office(repo, book, "real.xlsx", "xlsx")

    def _boom(file_path, file_name):
        raise RuntimeError("mineru 502")
    monkeypatch.setattr(repo.mineru_client, "parse_with_images", _boom)

    repo.process_source(sid)
    detail = repo.get_source(sid)
    assert detail.parse_status == "extracted"
    assert detail.parse_quality_warning is False              # 全保真兜底不是降级
    assert not (detail.error_message or "").startswith("[pdf-python-fallback]")
    elements = repo.source_elements(sid)
    assert all(e.metadata.get("parser") == "xlsx" for e in elements)
    assert any("Hello from xlsx fallback." in e.text for e in elements)


def test_cloud_only_config_does_not_upload_non_mineru_suffixes(tmp_path, monkeypatch):
    """云端上传只对 MinerU 真能解析的后缀发起：.csv 直接本地解析，绝不外发。"""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.delenv("MINERU_MODE", raising=False)          # 本地 off
    monkeypatch.delenv("MINERU_API_URL", raising=False)
    monkeypatch.setenv("MINERU_API_TOKEN", "tok")             # 云端 configured
    repo = _repository(Settings())
    csv_path = tmp_path / "rows.csv"
    csv_path.write_text("a,b\n1,2\n", encoding="utf-8")
    nb, sid = _seed_queued_office(repo, csv_path, "rows.csv", "csv")

    calls = []
    monkeypatch.setattr(
        repo.mineru_cloud_client, "parse_file_with_images",
        lambda path, data_id="": calls.append(path) or ([], {}),
    )

    repo.process_source(sid)

    assert calls == []                                        # 零次云端往返
    detail = repo.get_source(sid)
    assert detail.parse_status == "extracted"
    assert detail.parse_quality_warning is False              # 无 MinerU 路径可降级
    elements = repo.source_elements(sid)
    assert any(e.metadata.get("parser") == "csv" for e in elements)


def test_cloud_only_config_still_uploads_xlsx(tmp_path, monkeypatch):
    """反向护栏：xlsx 属于 MinerU 可解析后缀，云端分支照常发起。"""
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.delenv("MINERU_MODE", raising=False)
    monkeypatch.delenv("MINERU_API_URL", raising=False)
    monkeypatch.setenv("MINERU_API_TOKEN", "tok")
    repo = _repository(Settings())
    book = tmp_path / "book.xlsx"
    book.write_bytes(b"PK\x03\x04stub")
    nb, sid = _seed_queued_office(repo, book, "book.xlsx", "xlsx")

    calls = []
    monkeypatch.setattr(
        repo.mineru_cloud_client, "parse_file_with_images",
        lambda path, data_id="": (
            calls.append(path)
            or ([{"type": "text", "text": "From cloud.", "page_idx": 0}], {})
        ),
    )

    repo.process_source(sid)

    assert calls == [str(book)]
    assert repo.get_source(sid).parse_status == "extracted"
    elements = repo.source_elements(sid)
    assert any(e.metadata.get("parser") == "mineru" for e in elements)
    # 云端分支必须按后缀取标签前缀：默认的 "PDF" 会让 cloud-only 部署上传的工作簿
    # 拿到 "PDF p.1 ..." 与 source_format="pdf"，来源详情与引用当场答错格式。
    assert all(e.location_label.startswith("XLSX p.") for e in elements)
    assert {e.metadata.get("source_format") for e in elements} == {"xlsx"}


def _write_xlsx_rows(path, rows, cols=2):
    """行/列可控的真工作簿——云端对账的两个分母分别是 rows 与 rows*cols。"""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    for index in range(rows):
        sheet.append([f"r{index}c{col}" for col in range(cols)])
    book.save(str(path))
    return path


def _cloud_only_repo(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 't.db'}")
    monkeypatch.setenv("SILICON_NOTEBOOK_STORAGE_DIR", str(tmp_path / "storage"))
    monkeypatch.setenv("EVENT_LOG_ENABLED", "false")
    monkeypatch.setenv("LLM_LOG_ENABLED", "false")
    monkeypatch.delenv("MINERU_MODE", raising=False)           # 本地 off
    monkeypatch.delenv("MINERU_API_URL", raising=False)
    monkeypatch.setenv("MINERU_API_TOKEN", "tok")              # 云端 configured
    return _repository(Settings())


def test_cloud_workbook_output_must_clear_the_coverage_reconciliation(tmp_path, monkeypatch):
    """cloud-only 部署走的是摄取层自己的映射路径，不经 parse_xlsx——不在这里对账
    就等于给云端产出开了个绕过护栏的后门：一份丢掉大半行的表照样被采信、记成
    parser=mineru，静默数据损失。

    顺带钉住「拒收零持久化」：被丢弃的产出不能在资产表/磁盘上留下孤儿图片。
    """
    repo = _cloud_only_repo(tmp_path, monkeypatch)
    book = _write_xlsx_rows(tmp_path / "book.xlsx", 10)         # 10 行 × 2 列
    nb, sid = _seed_queued_office(repo, book, "book.xlsx", "xlsx")

    dropped = "".join(f"<tr><td>a{i}</td><td>b{i}</td></tr>" for i in range(3))
    monkeypatch.setattr(
        repo.mineru_cloud_client, "parse_file_with_images",
        lambda path, data_id="": (
            [
                {"type": "table", "table_body": f"<table>{dropped}</table>", "page_idx": 0},
                {"type": "image", "img_path": "images/a.png", "page_idx": 0},
            ],
            {"a.png": b"\x89PNG\r\n\x1a\n" + b"0" * 32},
        ),
    )

    repo.process_source(sid)

    assert repo.get_source(sid).parse_status == "extracted"
    elements = repo.source_elements(sid)
    assert all(e.metadata.get("parser") == "xlsx" for e in elements)  # 回落 openpyxl
    assert len(elements) == 10                                        # 全保真：一行不少
    assert repo.source_asset_ids(sid) == []                           # 零孤儿资产


def test_cloud_workbook_output_that_covers_the_workbook_is_kept(tmp_path, monkeypatch):
    """反向护栏：覆盖足额的云端产出照常采信，图片也照常持久化——两段式映射
    （探针不带 persist_image，通过才重映射）不能把图片一起弄丢。"""
    repo = _cloud_only_repo(tmp_path, monkeypatch)
    book = _write_xlsx_rows(tmp_path / "book.xlsx", 10)
    nb, sid = _seed_queued_office(repo, book, "book.xlsx", "xlsx")

    full = "".join(f"<tr><td>a{i}</td><td>b{i}</td></tr>" for i in range(10))
    png_bytes = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    monkeypatch.setattr(
        repo.mineru_cloud_client, "parse_file_with_images",
        lambda path, data_id="": (
            [
                {"type": "table", "table_body": f"<table>{full}</table>", "page_idx": 0},
                {"type": "image", "img_path": "images/a.png", "page_idx": 0},
            ],
            {"a.png": png_bytes},
        ),
    )

    repo.process_source(sid)

    assert repo.get_source(sid).parse_status == "extracted"
    elements = repo.source_elements(sid)
    assert all(e.metadata.get("parser") == "mineru" for e in elements)
    asset_ids = repo.source_asset_ids(sid)
    assert len(asset_ids) == 1
    image_el = next(e for e in elements if e.element_type == "image")
    assert image_el.metadata.get("asset_id") == asset_ids[0]
