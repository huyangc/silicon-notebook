"""Office Math (OMML) in the local DOCX / PPTX fallback parsers.

Before this, mammoth dropped every `m:oMath`, python-docx / python-pptx `.text`
skipped them, and python-pptx never yielded PowerPoint's equation boxes at all
(they live in `mc:AlternateContent`), losing the text around them too. The
contract now follows MinerU's: display math → its own `formula` element with
bare LaTeX and `text_format: latex`; inline math → `$latex$` in the text.
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest

from app.services.omml import (
    docx_math_package,
    drawingml_inline_text,
    inline_math_for,
    omml_display_lines,
    omml_to_latex,
    wordprocessing_paragraph_text,
)
from app.services.parsers import parse_builtin_source_file, parse_docx_basic, parse_pptx_basic
from tests.omml_testkit import (
    A14_NS,
    A_NS,
    M_NS,
    acc,
    bar,
    delim,
    docx_with,
    e,
    eq_arr,
    frac,
    func,
    group_chr,
    lim_low,
    lim_upp,
    matrix,
    nary,
    omath,
    omath_para,
    pptx_paragraph,
    pptx_with,
    r,
    rad,
    ssub,
    ssubsup,
    ssup,
    wrap,
)


def _latex(*children: str) -> str:
    return omml_to_latex(ElementTree.fromstring(omath(*children)))


def _rows(elements):
    return [(el.element_type, el.location_label, el.text, el.metadata) for el in elements]


# --- converter: one assertion per construct ---------------------------------

CONVERTER_CASES = [
    # fractions
    ("frac_bar", [frac(r("a"), r("b"))], r"\frac{a}{b}"),
    ("frac_linear", [frac(r("a"), r("b"), kind="lin")], "a/b"),
    ("frac_skewed", [frac(r("x+1"), r("2"), kind="skw")], "{x+1}/2"),
    ("frac_no_bar", [frac(r("n"), r("k"), kind="noBar")], r"\genfrac{}{}{0pt}{}{n}{k}"),
    # scripts
    ("superscript", [ssup(r("x"), r("2"))], "x^{2}"),
    ("subscript", [ssub(r("x"), r("i"))], "x_{i}"),
    ("sub_superscript", [ssubsup(r("x"), r("i"), r("2"))], "x_{i}^{2}"),
    ("superscript_compound_base", [ssup(r("x+y"), r("n"))], "{x+y}^{n}"),
    ("pre_script", [wrap("sPre", wrap("sub", r("a")), wrap("sup", r("b")), e(r("X")))], "{}_{a}^{b}X"),
    # radicals
    ("square_root", [rad(r("x"))], r"\sqrt{x}"),
    ("nth_root", [rad(r("x"), deg=r("3"))], r"\sqrt[3]{x}"),
    # n-ary
    ("sum", [nary("∑", r("i=1"), r("n"), ssub(r("x"), r("i")))], r"\sum_{i=1}^{n} x_{i}"),
    ("integral_default_chr", [nary(None, r("0"), r("1"), r("f"))], r"\int_{0}^{1} f"),
    ("product", [nary("∏", r("k"), "", r("a"))], r"\prod_{k} a"),
    ("integral_under_over", [nary("∫", r("a"), r("b"), r("g"), lim_loc="undOvr")], r"\int\limits_{a}^{b} g"),
    ("sum_sub_sup", [nary("∑", r("i"), "", r("y"), lim_loc="subSup")], r"\sum\nolimits_{i} y"),
    ("contour_integral", [nary("∮", r("C"), "", r("F"))], r"\oint_{C} F"),
    # delimiters
    ("parens", [delim(r("x"))], r"\left( x \right)"),
    ("brackets_separator", [delim(r("a"), r("b"), beg="[", end="]")], r"\left[ a \middle| b \right]"),
    ("braces_comma_sep", [delim(r("1"), r("2"), beg="{", end="}", sep=",")], r"\left\{ 1 , 2 \right\}"),
    ("open_only", [delim(r("x"), beg="", end="|")], r"\left. x \right|"),
    ("angle", [delim(r("ψ"), beg="⟨", end="⟩")], r"\left\langle \psi \right\rangle"),
    ("unknown_bracket_kept_as_text", [delim(r("x"), beg="⟦", end="⟧")], "⟦x⟧"),
    # matrices / arrays
    ("matrix", [matrix((r("a"), r("b")), (r("c"), r("d")))], r"\begin{matrix} a & b \\ c & d \end{matrix}"),
    (
        "matrix_in_parens",
        [delim(matrix((r("1"), r("0")), (r("0"), r("1"))))],
        r"\left( \begin{matrix} 1 & 0 \\ 0 & 1 \end{matrix} \right)",
    ),
    ("eq_array_aligned", [eq_arr(r("x&=1"), r("y&=2"))], r"\begin{aligned} x&=1 \\ y&=2 \end{aligned}"),
    ("eq_array_gathered", [eq_arr(r("a"), r("b"))], r"\begin{gathered} a \\ b \end{gathered}"),
    # functions
    ("function_sin", [func(r("sin", sty="p"), r("x"))], r"\sin{x}"),
    ("function_unknown_name", [func(r("sinc"), r("x"))], r"\operatorname{sinc}{x}"),
    (
        "limit",
        [func(lim_low(r("lim", sty="p"), r("n→∞")), ssub(r("a"), r("n")))],
        r"\lim_{n\rightarrow\infty}{a_{n}}",
    ),
    ("lim_low_generic", [lim_low(r("x"), r("k"))], r"\underset{k}{x}"),
    ("lim_upp", [lim_upp(r("="), r("def"))], r"\overset{def}{=}"),
    # accents / bars / group characters
    ("accent_default_hat", [acc(r("x"))], r"\hat{x}"),
    ("accent_vector", [acc(r("v"), "\u20d7")], r"\vec{v}"),
    ("accent_tilde", [acc(r("n"), "\u0303")], r"\tilde{n}"),
    ("accent_dot", [acc(r("x"), "\u0307")], r"\dot{x}"),
    ("accent_unknown", [acc(r("x"), "★")], r"\overset{★}{x}"),
    ("bar_bottom_default", [bar(r("x"))], r"\underline{x}"),
    ("bar_top", [bar(r("x"), pos="top")], r"\overline{x}"),
    ("group_underbrace_default", [group_chr(r("a+b"))], r"\underbrace{a+b}"),
    ("group_overbrace", [group_chr(r("a+b"), "⏞", "top")], r"\overbrace{a+b}"),
    ("group_arrow_top", [group_chr(r("x"), "→", "top")], r"\overset{\rightarrow}{x}"),
    ("border_box", [wrap("borderBox", e(r("E")))], r"\boxed{E}"),
    ("box_passthrough", [wrap("box", e(r("dx")))], "dx"),
    ("phantom_hidden", [wrap("phant", '<m:phantPr><m:show m:val="0"/></m:phantPr>', e(r("x")))], r"\phantom{x}"),
    # text and symbols
    ("greek_and_relations", [r("α≤β")], r"\alpha\le\beta"),
    ("command_then_letter_gets_space", [r("αx")], r"\alpha x"),
    ("operators", [r("a±b×c≠d")], r"a\pm b\times c\ne d"),
    ("calculus_symbols", [r("∂f/∂x→∞")], r"\partial f/\partial x\rightarrow\infty"),
    ("latex_specials_escaped", [r("$5 50% a_b {x} #")], r"\$5 50\% a\_b \{x\} \#"),
    ("ampersand_outside_array", [r("A&B")], r"A\&B"),
    ("math_italic_letters_fold", [r("𝑥+𝛼")], r"x+\alpha"),
    ("invisible_function_application_dropped", [r("f\u2061x")], "fx"),
    ("unicode_minus", [r("a−b")], "a-b"),
    ("normal_text_run", [r("if x>0", nor=True)], r"\text{if x>0}"),
    ("cjk_text_run", [r("速度", nor=False)], r"\text{速度}"),
    ("plain_style_function_name", [r("max", sty="p")], r"\max"),
    ("plain_style_word", [r("abc", sty="p")], r"\mathrm{abc}"),
    ("double_struck", [r("R", scr="double-struck")], r"\mathbb{R}"),
    ("bold", [r("v", sty="b")], r"\mathbf{v}"),
    # degradation: unknown constructs keep their text
    ("unknown_element_keeps_text", [wrap("futureThing", r("q"), r("=1"))], "q=1"),
    ("unknown_element_wraps_known", [wrap("futureThing", frac(r("1"), r("2")))], r"\frac{1}{2}"),
    # superscript-like symbols never stack on an existing superscript
    ("degree_after_superscript", [ssup(r("x"), r("2")), r("°")], r"x^{2}{}^{\circ}"),
    ("prime_after_superscript", [ssup(r("x"), r("2")), r("′")], r"x^{2}{}^{\prime}"),
    ("prime", [r("f′")], r"f^{\prime}"),
    ("double_prime_then_parens", [r("f″(x)")], r"f^{\prime\prime}(x)"),
    ("two_primes_in_a_row", [r("f′′")], r"f^{\prime}{}^{\prime}"),
    # combining marks typed after a base character
    ("combining_hat_on_previous", [r("x\u0302")], r"\hat{x}"),
    ("combining_underline", [r("a\u0332b")], r"\underline{a}b"),
    ("isolated_combining_dropped", [r("\u0302y")], "y"),
    ("unmapped_combining_dropped", [r("q\u0327")], "q"),
    # plain-text delimiter fallback uses a plain separator (no \middle)
    ("unknown_bracket_plain_separator", [delim(r("a"), r("b"), beg="«", end="»")], "«a | b»"),
    # a matrix's & inside an equation array is not an alignment point
    (
        "eq_array_matrix_ampersand_is_text",
        [eq_arr(matrix((r("a&b"), r("c"))))],
        r"\begin{gathered} \begin{matrix} a\&b & c \end{matrix} \end{gathered}",
    ),
    (
        "eq_array_alignment_beside_matrix",
        [eq_arr(r("M&="), matrix((r("1"),)))],
        r"\begin{aligned} M&= \\ \begin{matrix} 1 \end{matrix} \end{aligned}",
    ),
]


@pytest.mark.parametrize(
    ("children", "expected"),
    [case[1:] for case in CONVERTER_CASES],
    ids=[case[0] for case in CONVERTER_CASES],
)
def test_omml_construct_converts_to_latex(children, expected):
    assert _latex(*children) == expected


def test_converter_output_never_breaks_inline_dollar_delimiters():
    latex = _latex(r("$"), ssup(r("$"), r("2")))
    assert "$" not in latex.replace(r"\$", "")
    assert inline_math_for(ElementTree.fromstring(omath(r("$")))) == r"$\$$"


def test_tracked_deletions_inside_math_are_not_text():
    xml = omath(
        r("a"),
        f'<w:del xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">{r("z")}</w:del>',
        f'<w:ins xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">{r("b")}</w:ins>',
    )
    assert omml_to_latex(ElementTree.fromstring(xml)) == "ab"


def test_pathological_nesting_degrades_to_text_instead_of_raising():
    depth = 5000
    xml = omath("<m:box><m:e>" * depth + r("z") + "</m:e></m:box>" * depth)
    assert omml_to_latex(ElementTree.fromstring(xml)) == "z"


def test_math_paragraph_lines_and_inline_collapse():
    para = ElementTree.fromstring(omath_para(r("a=1"), r("b=2")))
    assert omml_display_lines(para) == ["a=1", "b=2"]
    assert inline_math_for(para) == r"$\begin{gathered} a=1 \\ b=2 \end{gathered}$"


def test_empty_math_yields_nothing():
    assert omml_to_latex(ElementTree.fromstring(omath())) == ""
    assert omml_display_lines(ElementTree.fromstring(omath_para(""))) == []


# --- DOCX (mammoth path) -----------------------------------------------------


def _parse_docx(path: Path):
    return parse_builtin_source_file("s1", str(path), path.name)


def test_docx_inline_and_display_equations(tmp_path):
    path = docx_with(tmp_path / "math.docx", [
        ["Energy ", ("math", omath(r("E"), r("="), r("m"), ssup(r("c"), r("2")))), " holds."],
        [("math", omath_para(frac(r("a"), r("b"))))],
        ["Before ", ("math", omath_para(nary("∑", r("i=1"), r("n"), ssub(r("x"), r("i"))))), " after"],
    ])
    assert _rows(_parse_docx(path)) == [
        ("paragraph", "DOCX paragraph 1", "Energy $E=mc^{2}$ holds.", {"parser": "mammoth", "paragraph_index": 1}),
        ("formula", "DOCX formula 1", r"\frac{a}{b}",
         {"parser": "mammoth", "formula_index": 1, "text_format": "latex"}),
        ("paragraph", "DOCX paragraph 2", "Before", {"parser": "mammoth", "paragraph_index": 2}),
        ("formula", "DOCX formula 2", r"\sum_{i=1}^{n} x_{i}",
         {"parser": "mammoth", "formula_index": 2, "text_format": "latex"}),
        ("paragraph", "DOCX paragraph 3", "after", {"parser": "mammoth", "paragraph_index": 3}),
    ]


def test_docx_multi_line_display_block_is_one_formula_per_line(tmp_path):
    path = docx_with(tmp_path / "lines.docx", [[("math", omath_para(r("x=1"), r("y=2")))]])
    assert [(el.element_type, el.text) for el in _parse_docx(path)] == [("formula", "x=1"), ("formula", "y=2")]


def test_docx_heading_list_item_and_table_keep_math_inline(tmp_path):
    path = docx_with(tmp_path / "places.docx", [
        ("heading", "Chapter"),
        ("bullet", ["item ", ("math", omath_para(ssup(r("x"), r("2"))))]),
        ("table", [[["cell ", ("math", omath(r("a<b")))], ["plain"]]]),
    ])
    elements = _parse_docx(path)
    item = next(el for el in elements if el.text.startswith("- "))
    assert item.element_type == "paragraph"
    assert item.text == "- item $x^{2}$"
    table = next(el for el in elements if el.element_type == "table")
    assert table.text == "cell $a<b$ | plain"
    assert "cell $a&lt;b$" in table.metadata["table_html"]
    assert not any(el.element_type == "formula" for el in elements)


def test_docx_equation_with_unknown_element_degrades_to_its_text(tmp_path):
    path = docx_with(tmp_path / "unknown.docx", [
        ["See ", ("math", omath(wrap("futureThing", r("q"), r("=1")))), " now"],
    ])
    assert [el.text for el in _parse_docx(path)] == ["See $q=1$ now"]


def test_docx_math_outside_a_paragraph_is_still_read(tmp_path):
    """`m:oMathPara` straight under `w:body` (some generators do this): the token
    run gets a paragraph of its own, or mammoth would never read it."""
    from docx import Document
    from docx.oxml import parse_xml

    document = Document()
    document.add_paragraph("lead")
    body = document.element.body
    body.insert(len(body) - 1, parse_xml(omath_para(rad(r("2")))))
    path = tmp_path / "body.docx"
    document.save(str(path))
    assert [(el.element_type, el.text) for el in _parse_docx(path)] == [
        ("paragraph", "lead"),
        ("formula", r"\sqrt{2}"),
    ]


def _add_footnote(path: Path, footnote_body: str) -> None:
    """Give the DOCX at `path` footnote 1 (python-docx cannot write footnotes)."""
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(path) as archive:
        members = {info.filename: archive.read(info) for info in archive.infolist()}
    members["word/footnotes.xml"] = (
        f'<w:footnotes xmlns:w="{w_ns}" xmlns:m="{M_NS}"><w:footnote w:id="1">'
        f"<w:p>{footnote_body}</w:p></w:footnote></w:footnotes>"
    ).encode()
    members["word/_rels/document.xml.rels"] = members["word/_rels/document.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        b'relationships/footnotes" Target="footnotes.xml"/></Relationships>',
    )
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/word/footnotes.xml" ContentType="application/'
        b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    )
    members["word/document.xml"] = members["word/document.xml"].replace(
        b"<w:t>See note</w:t></w:r>",
        b'<w:t>See note</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r>',
    )
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def test_docx_footnote_math_is_converted(tmp_path):
    path = docx_with(tmp_path / "footnote.docx", [["See note"], [("math", omath_para(frac(r("p"), r("q"))))]])
    _add_footnote(path, '<w:r><w:t xml:space="preserve">Note </w:t></w:r>' + omath(ssup(r("x"), r("2"))))
    texts = [(el.element_type, el.text) for el in _parse_docx(path)]
    assert ("formula", r"\frac{p}{q}") in texts
    assert any(kind == "paragraph" and "Note $x^{2}$" in text for kind, text in texts)


def test_docx_text_that_looks_like_a_token_is_left_alone(tmp_path):
    """The sentinel is chosen absent from every text part, so private-use text
    shaped like a token can never be mistaken for one."""
    first_private_use = "\ue000"
    lookalike = f"{first_private_use}0{first_private_use}"
    path = docx_with(tmp_path / "pua.docx", [[f"odd {lookalike} text"], ["v ", ("math", omath(r("y")))]])
    assert [el.text for el in _parse_docx(path)] == [f"odd {lookalike} text", "v $y$"]


def test_docx_python_docx_last_resort_keeps_math_inline(tmp_path, monkeypatch):
    import mammoth

    path = docx_with(tmp_path / "math.docx", [
        ["Energy ", ("math", omath(r("E"), r("="), r("m"), ssup(r("c"), r("2")))), " holds."],
        [("math", omath_para(frac(r("a"), r("b"))))],
        ("table", [[["v=", ("math", omath(rad(r("x"))))], ["w"]]]),
    ])

    def _boom(*args, **kwargs):
        raise RuntimeError("mammoth boom")

    monkeypatch.setattr(mammoth, "convert_to_html", _boom)
    assert _rows(_parse_docx(path)) == [
        ("paragraph", "DOCX paragraph 1", "Energy $E=mc^{2}$ holds.", {"parser": "docx", "paragraph_index": 1}),
        ("paragraph", "DOCX paragraph 2", r"$\frac{a}{b}$", {"parser": "docx", "paragraph_index": 2}),
        ("table_row", "DOCX table 1 row 1", r"v=$\sqrt{x}$ | w",
         {"parser": "docx", "table_index": 1, "row_index": 1}),
    ]


def test_docx_math_rewrite_failure_keeps_the_rest_of_the_document(tmp_path, monkeypatch):
    from app.services import parsers as parsers_module

    path = docx_with(tmp_path / "math.docx", [["Energy ", ("math", omath(r("E"))), " holds."]])

    def _boom(_path):
        raise RuntimeError("rewrite boom")

    monkeypatch.setattr(parsers_module, "docx_math_package", _boom)
    elements = _parse_docx(path)
    assert [(el.metadata["parser"], el.text) for el in elements] == [("mammoth", "Energy holds.")]


# --- DOCX: math-free documents are untouched ---------------------------------


def _rich_math_free_docx(path: Path) -> Path:
    return docx_with(path, [
        ("heading", "Pricing"),
        ["Costs $5 and $10, 50% off; a_b {x} & more."],
        ("bullet", ["first"]),
        ("table", [[["H1"], ["H2"]], [["v1"], ["v2"]]]),
        ["Last line."],
    ])


def test_math_free_docx_is_never_rewritten(tmp_path):
    path = _rich_math_free_docx(tmp_path / "plain.docx")
    with zipfile.ZipFile(path) as archive:
        # python-docx's template declares the math namespace like Word does;
        # the gate must not mistake the declaration for an equation.
        assert M_NS.encode() in archive.read("word/document.xml")
    assert docx_math_package(path) is None


def test_docx_mentioning_omath_in_prose_is_not_rewritten(tmp_path):
    path = docx_with(tmp_path / "prose.docx", [["The m:oMath element holds Office Math."]])
    assert docx_math_package(path) is None
    assert [el.text for el in _parse_docx(path)] == ["The m:oMath element holds Office Math."]


def test_math_free_docx_output_is_unchanged(tmp_path):
    path = _rich_math_free_docx(tmp_path / "plain.docx")
    assert _rows(_parse_docx(path)) == [
        ("heading", "DOCX heading 1", "Pricing", {"parser": "mammoth", "heading_index": 1, "heading_level": 1}),
        ("paragraph", "DOCX paragraph 1", "Costs $5 and $10, 50% off; a_b {x} & more.",
         {"parser": "mammoth", "paragraph_index": 1}),
        ("paragraph", "DOCX paragraph 2", "- first", {"parser": "mammoth", "paragraph_index": 2}),
        ("table", "DOCX table 1", "H1 | H2 ; v1 | v2",
         {"parser": "mammoth", "table_index": 1,
          "table_html": "<table><tr><td><p>H1</p></td><td><p>H2</p></td></tr>"
                        "<tr><td><p>v1</p></td><td><p>v2</p></td></tr></table>"}),
        ("paragraph", "DOCX paragraph 3", "Last line.", {"parser": "mammoth", "paragraph_index": 3}),
    ]


def test_math_free_docx_python_docx_output_is_unchanged(tmp_path):
    path = _rich_math_free_docx(tmp_path / "plain.docx")
    assert [el.text for el in parse_docx_basic("s1", path)] == [
        "Pricing", "Costs $5 and $10, 50% off; a_b {x} & more.", "first", "Last line.", "H1 | H2", "v1 | v2",
    ]


# --- PPTX (python-pptx path) -------------------------------------------------


def _parse_pptx(path: Path):
    return parse_builtin_source_file("s1", str(path), path.name)


def test_pptx_equation_boxes_become_formulas_and_text(tmp_path):
    path = pptx_with(tmp_path / "math.pptx", [[
        "Title box",
        ("equation", [pptx_paragraph(("math", omath_para(rad(r("x")))))]),
        ("equation", [pptx_paragraph("Inline ", ("math", omath(ssup(r("x"), r("2")))), " here")]),
        "Closing box",
    ]])
    assert _rows(_parse_pptx(path)) == [
        ("slide_text", "PPTX slide 1 shape 1", "Title box",
         {"parser": "python-pptx", "slide_number": 1, "shape_index": 1}),
        ("formula", "PPTX slide 1 formula 1", r"\sqrt{x}",
         {"parser": "python-pptx", "slide_number": 1, "formula_index": 1, "text_format": "latex"}),
        ("slide_text", "PPTX slide 1 shape 2", "Inline $x^{2}$ here",
         {"parser": "python-pptx", "slide_number": 1, "shape_index": 2}),
        ("slide_text", "PPTX slide 1 shape 3", "Closing box",
         {"parser": "python-pptx", "slide_number": 1, "shape_index": 3}),
    ]


def test_pptx_display_equation_splits_the_text_around_it(tmp_path):
    path = pptx_with(tmp_path / "split.pptx", [[
        ("equation", [
            pptx_paragraph("Given"),
            pptx_paragraph(("math", omath_para(frac(r("1"), r("2")), r("y=0")))),
            pptx_paragraph("then"),
        ]),
    ]])
    assert [(el.element_type, el.location_label, el.text) for el in _parse_pptx(path)] == [
        ("slide_text", "PPTX slide 1 shape 1", "Given"),
        ("formula", "PPTX slide 1 formula 1", r"\frac{1}{2}"),
        ("formula", "PPTX slide 1 formula 2", "y=0"),
        ("slide_text", "PPTX slide 1 shape 2", "then"),
    ]


def test_pptx_unwrapped_math_shape_and_grouped_equation(tmp_path):
    """Math in a plain `p:sp` (non-PowerPoint writers) and an equation box
    inside a group are both read."""
    from pptx.oxml import parse_xml

    from tests.omml_testkit import equation_shape_xml

    path = pptx_with(tmp_path / "plain.pptx", [[
        ("plain_math", [pptx_paragraph("k = ", ("math", omath(frac(r("1"), r("4")))))]),
    ]])
    from pptx import Presentation

    presentation = Presentation(str(path))
    slide = presentation.slides[0]
    group = slide.shapes.add_group_shape([])
    group.element.append(parse_xml(equation_shape_xml(300, pptx_paragraph(("math", omath_para(r("z")))))))
    presentation.save(str(path))

    assert [(el.element_type, el.text) for el in _parse_pptx(path)] == [
        ("slide_text", r"k = $\frac{1}{4}$"),
        ("formula", "z"),
    ]


def test_pptx_table_cell_and_notes_math_stay_inline(tmp_path):
    from pptx import Presentation
    from pptx.oxml import parse_xml

    path = pptx_with(tmp_path / "cells.pptx", [[
        ("table", [["H1", "H2"], ["v1", "v2"]]),
        ("notes_math", pptx_paragraph("Note ", ("math", omath_para(ssub(r("a"), r("0")))))),
    ]])
    presentation = Presentation(str(path))
    table = next(shape for shape in presentation.slides[0].shapes if shape.has_table).table
    paragraph = table.cell(1, 1)._tc.txBody.findall(f"{{{A_NS}}}p")[0]
    paragraph.append(parse_xml(f'<a14:m xmlns:a14="{A14_NS}">{omath(r("x<y"))}</a14:m>'))
    presentation.save(str(path))

    elements = _parse_pptx(path)
    table_element = next(el for el in elements if el.element_type == "table")
    assert table_element.text == "H1 | H2 ; v1 | v2$x<y$"
    assert "<td>v2$x&lt;y$</td>" in table_element.metadata["table_html"]
    notes = next(el for el in elements if el.element_type == "speaker_notes")
    assert notes.text == "Note $a_{0}$"


def test_pptx_non_math_alternate_content_is_still_not_read(tmp_path):
    """Only equation alternates are opened; any other AlternateContent keeps the
    behaviour it always had (python-pptx's shape tree does not yield it)."""
    from pptx import Presentation
    from pptx.oxml import parse_xml

    from tests.omml_testkit import MC_NS, _text_box_xml

    path = pptx_with(tmp_path / "alt.pptx", [[
        "Visible",
        ("equation", [pptx_paragraph(("math", omath_para(r("q"))))]),
    ]])
    presentation = Presentation(str(path))
    slide = presentation.slides[0]
    slide.shapes._spTree.append(parse_xml(
        f'<mc:AlternateContent xmlns:mc="{MC_NS}"><mc:Choice Requires="p14">'
        f'{_text_box_xml(400, pptx_paragraph("Ink choice"))}</mc:Choice>'
        f'<mc:Fallback>{_text_box_xml(400, pptx_paragraph("Ink fallback"))}</mc:Fallback></mc:AlternateContent>'
    ))
    presentation.save(str(path))
    assert [el.text for el in _parse_pptx(path)] == ["Visible", "q"]


def test_pptx_raw_xml_last_resort_keeps_math_inline_without_duplicates(tmp_path):
    path = pptx_with(tmp_path / "math.pptx", [[
        "Title box",
        ("equation", [pptx_paragraph(("math", omath_para(rad(r("x")))))]),
        ("equation", [pptx_paragraph("Inline ", ("math", omath(ssup(r("x"), r("2")))), " here")]),
    ]])
    assert [(el.location_label, el.text) for el in parse_pptx_basic("s1", path)] == [
        ("PPTX slide 1 shape 1", "Title box"),
        ("PPTX slide 1 shape 2", r"$\sqrt{x}$"),
        ("PPTX slide 1 shape 3", "Inline $x^{2}$ here"),
    ]


def test_pptx_raw_xml_reads_an_equation_alternate_once(tmp_path):
    """`.//p:sp` sees both halves of the AlternateContent; a Fallback that
    carries a text copy of the equation must not add a second element."""
    from pptx import Presentation
    from pptx.oxml import parse_xml

    from tests.omml_testkit import MC_NS, _text_box_xml

    path = pptx_with(tmp_path / "dup.pptx", [["Title box"]])
    presentation = Presentation(str(path))
    presentation.slides[0].shapes._spTree.append(parse_xml(
        f'<mc:AlternateContent xmlns:mc="{MC_NS}" xmlns:a14="{A14_NS}"><mc:Choice Requires="a14">'
        f'{_text_box_xml(500, pptx_paragraph("k=", ("math", omath(r("2")))))}</mc:Choice>'
        f'<mc:Fallback>{_text_box_xml(500, pptx_paragraph("k=2 (picture copy)"))}</mc:Fallback>'
        "</mc:AlternateContent>"
    ))
    presentation.save(str(path))
    assert [el.text for el in parse_pptx_basic("s1", path)] == ["Title box", "k=$2$"]


def test_math_free_pptx_output_is_unchanged(tmp_path):
    path = pptx_with(tmp_path / "plain.pptx", [
        ["Costs $5 and $10", ("table", [["H1", "H2"], ["v1", "v2"]]), ("notes", "Speaker $note$")],
        ["Second slide"],
    ])
    assert _rows(_parse_pptx(path)) == [
        ("slide_text", "PPTX slide 1 shape 1", "Costs $5 and $10",
         {"parser": "python-pptx", "slide_number": 1, "shape_index": 1}),
        ("table", "PPTX slide 1 table 1", "H1 | H2 ; v1 | v2",
         {"parser": "python-pptx", "slide_number": 1, "table_index": 1,
          "table_html": "<table><tr><td>H1</td><td>H2</td></tr><tr><td>v1</td><td>v2</td></tr></table>"}),
        ("speaker_notes", "PPTX slide 1 notes", "Speaker $note$",
         {"parser": "python-pptx", "slide_number": 1, "kind": "notes"}),
        ("slide_text", "PPTX slide 2 shape 1", "Second slide",
         {"parser": "python-pptx", "slide_number": 2, "shape_index": 1}),
    ]


# --- review round 1 -----------------------------------------------------------


def _tiny_png() -> bytes:
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(b"\0\0\0\0")) + chunk(b"IEND", b"")


def test_docx_math_package_never_reads_image_payloads(tmp_path):
    """mammoth never opens image bytes (`convert_image` ignores them), so the
    rewritten package must not either: a corrupt member would raise on read and
    a zip-bomb member would be inflated into memory / the temp file."""
    import tracemalloc

    path = docx_with(tmp_path / "payload.docx", [["a ", ("math", omath(ssup(r("x"), r("2"))))]])
    marker = b"CORRUPT-PAYLOAD-"
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr(zipfile.ZipInfo("word/media/corrupt.bin"), marker * 64)
        archive.writestr("word/media/bomb.bin", b"\0" * (64 * 1024 * 1024), compress_type=zipfile.ZIP_DEFLATED)
    raw = bytearray(path.read_bytes())
    raw[raw.index(marker)] = ord("X")
    path.write_bytes(bytes(raw))
    with zipfile.ZipFile(path) as archive, pytest.raises(zipfile.BadZipFile):
        archive.read("word/media/corrupt.bin")  # premise: reading it fails

    tracemalloc.start()
    try:
        docx_math = docx_math_package(path)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert docx_math is not None
    try:
        with zipfile.ZipFile(docx_math.package) as package:
            assert package.getinfo("word/media/bomb.bin").file_size == 0
            assert package.getinfo("word/media/corrupt.bin").file_size == 0
            assert package.getinfo("word/styles.xml").file_size > 0
    finally:
        docx_math.close()
    assert peak < 16 * 1024 * 1024, peak
    assert [el.text for el in _parse_docx(path)] == ["a $x^{2}$"]


def test_docx_picture_reads_the_same_with_math_present(tmp_path):
    import io

    from docx import Document
    from docx.oxml import parse_xml

    def build(name: str, with_math: bool) -> Path:
        document = Document()
        document.add_paragraph("before")
        run = document.add_paragraph().add_run()
        run.add_picture(io.BytesIO(_tiny_png()))
        run._r.xpath(".//wp:docPr")[0].set("descr", "diagram alt")
        document.add_paragraph("after")
        if with_math:
            document.add_paragraph()._p.append(parse_xml(omath_para(r("z"))))
        path = tmp_path / name
        document.save(str(path))
        return path

    plain = _rows(_parse_docx(build("plain.docx", False)))
    with_math = _rows(_parse_docx(build("math.docx", True)))
    assert any("diagram alt" in row[2] for row in plain)
    assert with_math == plain + [
        ("formula", "DOCX formula 1", "z", {"parser": "mammoth", "formula_index": 1, "text_format": "latex"})
    ]


def test_pick_sentinel_is_one_linear_pass_and_sees_every_reference_spelling():
    import time

    from app.services.omml import _pick_sentinel

    assert _pick_sentinel([b"&#x0E000; &#057345;"]) == "\ue002"
    every = "".join(chr(code) for code in range(0xE000, 0xF900)).encode()
    start = time.perf_counter()
    assert _pick_sentinel([b"x" * (5 * 1024 * 1024) + every]) is None
    assert time.perf_counter() - start < 2.0


def test_docx_without_a_free_sentinel_keeps_formulas_as_inline_text(tmp_path):
    every = "".join(chr(code) for code in range(0xE000, 0xF900))
    path = docx_with(tmp_path / "full.docx", [
        [every],
        ["v ", ("math", omath(r("y")))],
        [("math", omath_para(frac(r("1"), r("2"))))],
    ])
    assert [el.text for el in _parse_docx(path)] == [every, "v $y$", r"$\frac{1}{2}$"]


def test_adjacent_formulas_never_fuse_into_display_markers(tmp_path):
    parts = [("math", omath(r("a"))), ("math", omath(r("b"))), "$5"]
    docx = docx_with(tmp_path / "adjacent.docx", [parts])
    assert [el.text for el in _parse_docx(docx)] == ["$a$ $b$ $5"]
    assert [el.text for el in parse_docx_basic("s1", docx)] == ["$a$ $b$ $5"]
    pptx = pptx_with(tmp_path / "adjacent.pptx", [[
        ("equation", [pptx_paragraph(("math", omath(r("a"))), ("math", omath(r("b"))), "$5")]),
        ("notes_math", pptx_paragraph("N ", ("math", omath_para(r("a"), r("b"))))),
    ]])
    elements = _parse_pptx(pptx)
    assert [el.text for el in elements if el.element_type == "slide_text"] == ["$a$ $b$ $5"]
    assert [el.text for el in elements if el.element_type == "speaker_notes"] == [
        r"N $\begin{gathered} a \\ b \end{gathered}$"
    ]


def test_non_utf8_declared_parts_are_left_alone():
    from app.services.omml import _rewrite_docx_part

    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = f'<w:document xmlns:w="{w_ns}"><w:body><w:p>{omath(r("x"))}</w:p></w:body></w:document>'
    for declaration in ('<?xml version="1.0" encoding = "windows-1252"?>', "<?xml version='1.0' encoding='latin-1'?>"):
        assert _rewrite_docx_part((declaration + body).encode(), [], "\ue000") is None
    assert _rewrite_docx_part(('<?xml version="1.0" encoding = "UTF-8"?>' + body).encode(), [], "\ue000")


def test_deeply_wrapped_drawingml_runs_do_not_raise():
    depth = 3000
    xml = (
        f'<a:txBody xmlns:a="{A_NS}" xmlns:a14="{A14_NS}"><a:p>'
        + "<a:wrap>" * depth + "<a:r><a:t>deep</a:t></a:r>" + "</a:wrap>" * depth
        + f"<a14:m>{omath(r('q'))}</a14:m></a:p></a:txBody>"
    )
    assert drawingml_inline_text(ElementTree.fromstring(xml)) == "deep$q$"


def test_wordprocessing_text_matches_python_docx_on_breaks():
    from docx.oxml import parse_xml
    from docx.text.paragraph import Paragraph

    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    paragraph = parse_xml(
        f'<w:p xmlns:w="{w_ns}"><w:r><w:t>a</w:t><w:br w:type="page"/><w:t>b</w:t><w:br/>'
        '<w:t>c</w:t><w:br w:type="column"/><w:br w:type="textWrapping"/><w:cr/><w:tab/><w:noBreakHyphen/>'
        "</w:r></w:p>"
    )
    assert wordprocessing_paragraph_text(paragraph) == Paragraph(paragraph, None).text == "ab\nc\n\n\t-"


def test_math_free_documents_never_reach_the_math_machinery(tmp_path, monkeypatch):
    """Zero scan cost on math-free documents: past the byte / element gates,
    none of the math code may run."""
    from app.services import omml as omml_module
    from app.services import parsers as parsers_module

    def boom(*args, **kwargs):
        raise AssertionError("math machinery ran on a math-free document")

    for name in ("_DocxMathScanner", "_pick_sentinel", "_rewrite_docx_part"):
        monkeypatch.setattr(omml_module, name, boom)
    for name in (
        "_pptx_iter_shapes_with_math", "_pptx_math_units", "_pptx_math_notes_text",
        "_pptx_math_cell_text", "_docx_paragraph_text", "_pptx_basic_math_text",
    ):
        monkeypatch.setattr(parsers_module, name, boom)

    docx = _rich_math_free_docx(tmp_path / "plain.docx")
    assert docx_math_package(docx) is None
    assert {el.metadata["parser"] for el in _parse_docx(docx)} == {"mammoth"}
    assert len(parse_docx_basic("s1", docx)) == 6
    pptx = pptx_with(tmp_path / "plain.pptx", [
        ["Costs $5", ("table", [["H1", "H2"], ["v1", "v2"]]), ("notes", "Speaker note")],
    ])
    assert [el.metadata["parser"] for el in _parse_pptx(pptx)] == ["python-pptx"] * 3
    # The raw-XML path must decide on the part bytes alone.
    monkeypatch.setattr(parsers_module, "contains_omml", boom)
    assert [el.text for el in parse_pptx_basic("s1", pptx)] == ["Costs $5", "Speaker note"]
