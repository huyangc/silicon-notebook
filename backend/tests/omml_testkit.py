"""Builders for DOCX / PPTX fixtures that carry Office Math (OMML).

Shared by the OMML parser tests and the before/after benchmark. The XML mirrors
what Word / PowerPoint write: Word puts `m:oMath` (inline) and `m:oMathPara`
(display) straight into `w:p`; PowerPoint wraps every equation text box in
`mc:AlternateContent` whose `mc:Choice Requires="a14"` holds the real shape
(`a:p > a14:m > m:oMathPara`) and whose `mc:Fallback` holds a picture-filled
copy with a single blank run.
"""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


# --- OMML snippets (prefix `m:`; the root helpers declare it) ---------------


def r(text: str, sty: str | None = None, nor: bool = False, scr: str | None = None) -> str:
    props = ""
    if sty:
        props += f'<m:sty m:val="{sty}"/>'
    if scr:
        props += f'<m:scr m:val="{scr}"/>'
    if nor:
        props += "<m:nor/>"
    rpr = f"<m:rPr>{props}</m:rPr>" if props else ""
    return f'<m:r>{rpr}<m:t xml:space="preserve">{escape(text)}</m:t></m:r>'


def wrap(tag: str, *children: str) -> str:
    return f"<m:{tag}>{''.join(children)}</m:{tag}>"


def e(*children: str) -> str:
    return wrap("e", *children)


def prop(name: str, value: str) -> str:
    return f'<m:{name} m:val="{escape(value, {chr(34): "&quot;"})}"/>'


def frac(num: str, den: str, kind: str | None = None) -> str:
    fpr = f"<m:fPr>{prop('type', kind)}</m:fPr>" if kind else ""
    return f"<m:f>{fpr}{wrap('num', num)}{wrap('den', den)}</m:f>"


def ssup(base: str, sup: str) -> str:
    return f"<m:sSup>{e(base)}{wrap('sup', sup)}</m:sSup>"


def ssub(base: str, sub: str) -> str:
    return f"<m:sSub>{e(base)}{wrap('sub', sub)}</m:sSub>"


def ssubsup(base: str, sub: str, sup: str) -> str:
    return f"<m:sSubSup>{e(base)}{wrap('sub', sub)}{wrap('sup', sup)}</m:sSubSup>"


def rad(body: str, deg: str = "") -> str:
    radpr = "" if deg else "<m:radPr><m:degHide m:val=\"1\"/></m:radPr>"
    return f"<m:rad>{radpr}{wrap('deg', deg)}{e(body)}</m:rad>"


def nary(char: str | None, sub: str, sup: str, body: str, lim_loc: str | None = None) -> str:
    props = (prop("chr", char) if char is not None else "") + (prop("limLoc", lim_loc) if lim_loc else "")
    pr = f"<m:naryPr>{props}</m:naryPr>" if props else ""
    return f"<m:nary>{pr}{wrap('sub', sub)}{wrap('sup', sup)}{e(body)}</m:nary>"


def delim(*items: str, beg: str | None = None, end: str | None = None, sep: str | None = None) -> str:
    props = ""
    if beg is not None:
        props += prop("begChr", beg)
    if sep is not None:
        props += prop("sepChr", sep)
    if end is not None:
        props += prop("endChr", end)
    pr = f"<m:dPr>{props}</m:dPr>" if props else ""
    return f"<m:d>{pr}{''.join(e(item) for item in items)}</m:d>"


def matrix(*rows: tuple[str, ...]) -> str:
    return "<m:m>" + "".join("<m:mr>" + "".join(e(cell) for cell in row) + "</m:mr>" for row in rows) + "</m:m>"


def func(name: str, body: str) -> str:
    return f"<m:func>{wrap('fName', name)}{e(body)}</m:func>"


def acc(body: str, char: str | None = None) -> str:
    pr = f"<m:accPr>{prop('chr', char)}</m:accPr>" if char is not None else ""
    return f"<m:acc>{pr}{e(body)}</m:acc>"


def bar(body: str, pos: str | None = None) -> str:
    pr = f"<m:barPr>{prop('pos', pos)}</m:barPr>" if pos else ""
    return f"<m:bar>{pr}{e(body)}</m:bar>"


def group_chr(body: str, char: str | None = None, pos: str | None = None) -> str:
    props = (prop("chr", char) if char is not None else "") + (prop("pos", pos) if pos else "")
    pr = f"<m:groupChrPr>{props}</m:groupChrPr>" if props else ""
    return f"<m:groupChr>{pr}{e(body)}</m:groupChr>"


def lim_low(body: str, lim: str) -> str:
    return f"<m:limLow>{e(body)}{wrap('lim', lim)}</m:limLow>"


def lim_upp(body: str, lim: str) -> str:
    return f"<m:limUpp>{e(body)}{wrap('lim', lim)}</m:limUpp>"


def eq_arr(*rows: str) -> str:
    return "<m:eqArr>" + "".join(e(row) for row in rows) + "</m:eqArr>"


def omath(*children: str) -> str:
    return f'<m:oMath xmlns:m="{M_NS}">{"".join(children)}</m:oMath>'


def omath_para(*lines: str) -> str:
    body = "".join(f"<m:oMath>{line}</m:oMath>" for line in lines)
    return f'<m:oMathPara xmlns:m="{M_NS}">{body}</m:oMathPara>'


# --- DOCX ------------------------------------------------------------------


def docx_with(path: Path, blocks: list) -> Path:
    """Build a DOCX from `blocks`. Each block is a list of parts for one
    paragraph: a str is a text run, ("math", xml) an OMML element appended in
    place. ("heading", text) / ("table", [[cell parts]]) make those blocks."""
    from docx import Document
    from docx.oxml import parse_xml

    document = Document()
    for block in blocks:
        if isinstance(block, tuple) and block[0] == "heading":
            document.add_heading(block[1], level=1)
            continue
        if isinstance(block, tuple) and block[0] == "table":
            rows = block[1]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for col_index, parts in enumerate(row):
                    paragraph = table.rows[row_index].cells[col_index].paragraphs[0]
                    _fill_paragraph(paragraph, parts, parse_xml)
            continue
        style = None
        if isinstance(block, tuple) and block[0] == "bullet":
            style, block = "List Bullet", block[1]
        paragraph = document.add_paragraph(style=style)
        _fill_paragraph(paragraph, block, parse_xml)
    document.save(str(path))
    return path


def _fill_paragraph(paragraph, parts, parse_xml) -> None:
    for part in parts:
        if isinstance(part, str):
            paragraph.add_run(part)
        else:
            paragraph._p.append(parse_xml(part[1]))


# --- PPTX ------------------------------------------------------------------


def _text_box_xml(shape_id: int, paragraphs_xml: str, fallback: bool = False) -> str:
    fill = '<a:blipFill><a:blip r:embed="rId99" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/></a:blipFill>' if fallback else ""
    return (
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}" xmlns:a14="{A14_NS}">'
        f'<p:nvSpPr><p:cNvPr id="{shape_id}" name="TextBox {shape_id}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{fill}</p:spPr>'
        f"<p:txBody><a:bodyPr/><a:lstStyle/>{paragraphs_xml}</p:txBody></p:sp>"
    )


def pptx_paragraph(*parts) -> str:
    """One `a:p`: str parts are runs, ("math", xml) parts go inside `a14:m`."""
    body = ""
    for part in parts:
        if isinstance(part, str):
            body += f'<a:r><a:rPr lang="en-US"/><a:t>{escape(part)}</a:t></a:r>'
        else:
            body += f"<a14:m>{part[1]}</a14:m>"
    return f"<a:p>{body}</a:p>"


def equation_shape_xml(shape_id: int, *paragraphs: str) -> str:
    """PowerPoint's equation text box: AlternateContent(Choice=a14 math, Fallback=picture)."""
    blank = '<a:p><a:r><a:rPr lang="en-US"><a:noFill/></a:rPr><a:t> </a:t></a:r></a:p>'
    return (
        f'<mc:AlternateContent xmlns:mc="{MC_NS}" xmlns:a14="{A14_NS}">'
        f'<mc:Choice Requires="a14">{_text_box_xml(shape_id, "".join(paragraphs))}</mc:Choice>'
        f"<mc:Fallback>{_text_box_xml(shape_id, blank, fallback=True)}</mc:Fallback>"
        "</mc:AlternateContent>"
    )


def pptx_with(path: Path, slides: list) -> Path:
    """Build a PPTX. Each slide is a list of items: a str adds a plain text box,
    ("equation", [paragraph xml...]) adds an AlternateContent equation box,
    ("plain_math", [paragraph xml...]) adds an unwrapped `p:sp` holding math,
    ("notes", text) sets the speaker notes, ("notes_math", paragraph xml)
    puts a paragraph with math in the notes body."""
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.util import Inches

    presentation = Presentation()
    shape_id = 100
    for items in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for item in items:
            shape_id += 1
            if isinstance(item, str):
                box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
                box.text_frame.text = item
            elif item[0] == "equation":
                slide.shapes._spTree.append(parse_xml(equation_shape_xml(shape_id, *item[1])))
            elif item[0] == "plain_math":
                slide.shapes._spTree.append(parse_xml(_text_box_xml(shape_id, "".join(item[1]))))
            elif item[0] == "table":
                rows = item[1]
                table = slide.shapes.add_table(
                    len(rows), len(rows[0]), Inches(1), Inches(3), Inches(4), Inches(1)
                ).table
                for row_index, row in enumerate(rows):
                    for col_index, value in enumerate(row):
                        table.cell(row_index, col_index).text = value
            elif item[0] == "notes":
                slide.notes_slide.notes_text_frame.text = item[1]
            elif item[0] == "notes_math":
                frame = slide.notes_slide.notes_text_frame
                frame.text = ""
                tx_body = frame._txBody
                for paragraph in list(tx_body.findall(f"{{{A_NS}}}p")):
                    tx_body.remove(paragraph)
                tx_body.append(parse_xml(item[1].replace("<a:p>", f'<a:p xmlns:a="{A_NS}" xmlns:a14="{A14_NS}">', 1)))
    presentation.save(str(path))
    return path
