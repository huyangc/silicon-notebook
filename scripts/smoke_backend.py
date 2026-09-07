from __future__ import annotations

import io
import json
import os
import tempfile
import zipfile
from pathlib import Path

from app.core.config import Settings
from app.models.schemas import (
    AskRequest,
    Evidence,
    FeedbackRequest,
    KnowledgeUpdate,
    MergeRequest,
    NotebookCreate,
    NotebookUpdate,
    SourceElement,
)
from app.services.parsers import mineru_content_list_to_elements
from app.services.repository import UploadedSourceFile
from app.services.retrieval import score_knowledge
from app.services.sqlite_repository import SQLiteRepository


def _offline_settings(root: Path, *, event_log_dir: Path | None = None) -> Settings:
    # llm_log_path 的目录必须与 event_log_dir 对齐(见 event_logging.llm_log_dir_aligned):
    # 否则 SQLiteRepository 构造时会打印「LLM_LOG_PATH 的目录与 EVENT_LOG_DIR 不一致」告警。
    # 默认 llm_log_path=.local/logs/llm.jsonl 指向仓库根,而这里 event_log_dir 指向本 check
    # 的临时目录,两者天然不一致——故把 llm_log_path 也钉进同一临时目录消除告警。
    # (这几个字段用 Field(default, env=...) 声明、按字段名绑定,Settings(llm_log_path=...)
    #  kwarg 生效;只有 validation_alias 声明的字段才须用别名名传参、否则被静默忽略。)
    _log_dir = event_log_dir or (root / "logs")
    return Settings(
        database_url=f"sqlite:///{root / 'silicon_notebook.db'}",
        storage_dir=str(root / "storage"),
        mineru_mode="off",
        model_services_config="",
        event_log_dir=str(_log_dir),
        llm_log_path=str(_log_dir / "llm.jsonl"),
        # P4: 摄取默认不抽 KG;冒烟脚本要端到端验证 KG 抽取/存储/检索全管线,
        # 故显式开启自动抽取(offline 下产出 'no-llm' extraction run,正是各断言所期)。
        kg_auto_extract=True,
    )


def _ev(element_id: str, quoted_span: str) -> Evidence:
    return Evidence(
        source_id="s",
        source_title="t",
        element_id=element_id,
        element_type="paragraph",
        location_label="p",
        quoted_span=quoted_span,
        confidence=0.7,
    )


def _source_evidence(
    repository: SQLiteRepository, source_id: str, quote: str | None = None
) -> dict:
    source = repository.get_source(source_id)
    element = next(el for el in repository.source_elements(source_id) if el.text.strip())
    span = quote or element.text.strip()[:300]
    return Evidence(
        source_id=source.id,
        source_title=source.title,
        element_id=element.id,
        element_type=element.element_type,
        location_label=element.location_label,
        quoted_span=span,
        confidence=1.0,
    ).model_dump()


def _insert_rule(
    repository: SQLiteRepository,
    notebook_id: str,
    source_id: str,
    *,
    title: str,
    statement: str,
    applies_to: list[str] | None = None,
    recommendation: str = "",
    risk_if_ignored: str = "",
) -> str:
    payload = {
        "title": title,
        "statement": statement,
        "applies_to": applies_to or [],
        "recommendation": recommendation,
        "risk_if_ignored": risk_if_ignored,
        "severity": "medium",
    }
    return repository.maintenance.seed_rule_object(
        notebook_id,
        payload=payload,
        evidence=[_source_evidence(repository, source_id)],
        source_id=source_id,
    )


def _latest_extraction_run(repository: SQLiteRepository, source_id: str) -> dict:
    row = repository.maintenance.latest_extraction_run(source_id)
    assert row is not None, f"source {source_id} should have an extraction run"
    return dict(row)


def check_extraction_profiles() -> None:
    """KG extraction profiles match the current supported document types."""
    from app.services.extraction_profiles import (
        OBJECT_SCHEMAS,
        PROFILES,
        detect_doc_type,
        resolve_profile,
    )

    def el(index: int, text: str, element_type: str = "paragraph") -> SourceElement:
        return SourceElement(
            id=f"e{index}",
            source_id="s",
            element_type=element_type,
            location_label=f"L{index}",
            text=text,
            metadata={},
        )

    assert set(PROFILES) == {"academic_paper", "textbook"}
    assert set(OBJECT_SCHEMAS) == {"concept", "claim", "formula", "procedure"}

    paper = [
        el(1, "Abstract: this paper studies conditional memory modules."),
        el(2, "We propose Engram and report experimental results."),
        el(3, "Related work compares the method with MoE."),
        el(4, "References: [1] Smith et al., 2021."),
    ]
    assert detect_doc_type("Engram paper", paper) == "academic_paper"
    assert resolve_profile("textbook", "Engram paper", paper).id == "academic_paper"

    textbook = [
        el(1, "Chapter 2 reviews transformer memory."),
        el(2, "Definition: a cache stores reusable activations."),
        el(3, "Exercise: derive the retrieval complexity."),
    ]
    assert detect_doc_type("Memory textbook", textbook) == "textbook"

    neutral = [el(1, "Neutral notes about analog layout spacing margins.")]
    assert detect_doc_type("Notes", neutral) is None
    assert resolve_profile("textbook", "Notes", neutral).id == "textbook"
    assert resolve_profile(None, "Notes", neutral).id == "academic_paper"

    for profile in PROFILES.values():
        assert profile.object_types == ["concept", "claim", "formula", "procedure"]
        for schema in profile.schemas():
            assert schema.primary == "name"
            assert "section_path" in schema.fields


def check_object_schemas() -> None:
    """Editable schema registry now seeds the four KG object types."""
    from app.models.schemas import ObjectSchemaCreate, ObjectSchemaUpdate

    with tempfile.TemporaryDirectory(prefix="sn-schema-") as tmp:
        root = Path(tmp)
        repo = SQLiteRepository(_offline_settings(root))
        seeded = {schema.object_type for schema in repo.list_object_schemas()}
        assert {"concept", "claim", "formula", "procedure"} <= seeded, seeded
        eff = repo.effective_schemas()
        assert eff["claim"].primary == "name"

        repo.create_object_schema(
            ObjectSchemaCreate(
                object_type="Process Window",
                fields=["name", "limit"],
                label="工艺窗口",
            )
        )
        custom = next(
            schema for schema in repo.list_object_schemas()
            if schema.object_type == "process_window"
        )
        assert custom.source == "custom" and custom.label == "工艺窗口"
        repo.update_object_schema(
            "process_window",
            ObjectSchemaUpdate(fields=["name", "limit", "margin"]),
        )
        assert "margin" in repo.effective_schemas()["process_window"].fields
        repo.update_object_schema("process_window", ObjectSchemaUpdate(status="disabled"))
        assert "process_window" not in repo.effective_schemas()
        repo.delete_object_schema("process_window")
        assert all(
            schema.object_type != "process_window"
            for schema in repo.list_object_schemas()
        )

        try:
            repo.delete_object_schema("concept")
            raise AssertionError("builtin schema should not be deletable")
        except ValueError:
            pass

        nb = repo.create_notebook(NotebookCreate(name="S", purpose="p", primary_domain="d"))
        assert repo.propose_schemas(nb.id) == []


def check_kg_record_binding() -> None:
    """KG nodes become product objects only when evidence can bind to source elements."""
    from app.services import kg_ingest
    from app.services.kg.models import Edge, Evidence as KGEvidence, KnowledgeGraph, Node

    elements = [
        SourceElement(
            id="el-1",
            source_id="src-1",
            element_type="paragraph",
            location_label="p.1",
            text="Engram is a conditional memory module for Transformer backbones.",
            metadata={},
        ),
        SourceElement(
            id="el-2",
            source_id="src-1",
            element_type="paragraph",
            location_label="p.2",
            text="The quiet ground return path avoids noise coupling in analog inputs.",
            metadata={},
        ),
    ]
    graph = KnowledgeGraph(
        doc_id="source.md",
        doc_type="academic",
        nodes=[
            Node(
                id="n1",
                type="Claim",
                name="Engram is a conditional memory module.",
                evidence=[
                    KGEvidence(
                        file="source.md",
                        char_start=0,
                        char_end=10,
                        line_start=1,
                        line_end=1,
                        quote="Engram is a conditional memory module",
                    )
                ],
            ),
            Node(
                id="n2",
                type="Concept",
                name="quiet ground return path",
                evidence=[
                    KGEvidence(
                        file="source.md",
                        char_start=0,
                        char_end=10,
                        line_start=2,
                        line_end=2,
                        quote="quiet ground return path avoids noise coupling",
                    )
                ],
            ),
            Node(
                id="n3",
                type="Claim",
                name="hallucinated unsupported claim",
                evidence=[
                    KGEvidence(
                        file="source.md",
                        char_start=0,
                        char_end=10,
                        line_start=3,
                        line_end=3,
                        quote="this text is not present",
                    )
                ],
            ),
        ],
        edges=[
            Edge(
                id="e1",
                type="about",
                source_id="n1",
                target_id="n2",
                evidence=[],
            ),
            Edge(
                id="e2",
                type="supports",
                source_id="n3",
                target_id="n1",
                evidence=[],
            ),
        ],
    )
    objects, relations = kg_ingest.build_records(graph, "src-1", "source.md", elements)
    assert {obj["local_id"] for obj in objects} == {"n1", "n2"}
    assert len(relations) == 1 and relations[0]["edge_type"] == "about"
    assert all(obj["evidence"] for obj in objects)


def check_kg_extract_window_grounding() -> None:
    """Window extraction drops invalid node types and ungrounded edge references."""
    from app.services.kg.extract import extract_window
    from app.services.kg.parsing import SourceElementQ

    class FakeClient:
        def chat_json(self, _messages, _schema_hint):
            return json.dumps(
                {
                    "nodes": [
                        {"local_id": "c", "type": "Concept", "name": "Engram", "ev": 0},
                        {
                            "local_id": "claim",
                            "type": "Claim",
                            "name": "Engram separates storage from computation",
                            "ev": 1,
                        },
                        {"local_id": "bad", "type": "Rule", "name": "invalid", "ev": 0},
                        {
                            "local_id": "lost",
                            "type": "Claim",
                            "name": "not in source",
                            "ev": 99,
                        },
                    ],
                    "edges": [
                        {
                            "type": "about",
                            "source": "claim",
                            "target": "c",
                            "ev": 1,
                        },
                        {
                            "type": "supports",
                            "source": "lost",
                            "target": "claim",
                            "ev": 1,
                        },
                    ],
                }
            )

    elements = [
        SourceElementQ(
            id="q-1",
            type="paragraph",
            text="Engram is a model component.",
            file="source.md",
            char_start=0,
            char_end=28,
            line_start=1,
            line_end=1,
        ),
        SourceElementQ(
            id="q-2",
            type="paragraph",
            text="Engram separates storage from computation.",
            file="source.md",
            char_start=29,
            char_end=72,
            line_start=2,
            line_end=2,
        ),
    ]
    nodes, edges = extract_window(FakeClient(), elements, "Intro", "academic", 0)
    assert {node.type for node in nodes} == {"Concept", "Claim"}
    assert len(edges) == 1 and edges[0].type == "about"


def check_kg_windowing() -> None:
    """Large markdown sources are split into multiple KG windows with full coverage."""
    from app.services.kg.windowing import windows_with_elements

    text = "# Chapter\n\n" + "\n\n".join(
        f"Paragraph {i}: {'x' * 700}" for i in range(12)
    )
    pairs = windows_with_elements(text, "large.md", None, n=1200, m=100)
    assert len(pairs) > 1
    covered = {element.text.split(":", 1)[0] for _, elements in pairs for element in elements}
    assert {f"Paragraph {i}" for i in range(12)} <= covered


def check_kg_store_ask_and_conversations() -> None:
    """KG storage, graph browse, Ask retrieval, citations, and conversations."""
    with tempfile.TemporaryDirectory(prefix="sn-kg-") as tmp:
        root = Path(tmp)
        repo = SQLiteRepository(_offline_settings(root))
        nb = repo.create_notebook(NotebookCreate(name="KG", purpose="p"))
        uploaded = repo.upload_sources(
            nb.id,
            [
                UploadedSourceFile(
                    file_name="source.md",
                    content_type="text/markdown",
                    content=(
                        "# Engram\n\n"
                        "Engram is a conditional memory module designed to augment "
                        "Transformer backbones by separating storage from computation.\n\n"
                        "The quiet ground return path avoids noise coupling in analog inputs.\n"
                    ).encode("utf-8"),
                )
            ],
        )
        source_id = uploaded[0].id
        run = _latest_extraction_run(repo, source_id)
        assert run["run_type"] == "kg"
        assert run["status"] == "completed"
        assert run["error_message"] == "no-llm"

        evidence = [_source_evidence(repo, source_id)]
        objects = [
            {
                "local_id": "claim-engram",
                "object_type": "claim",
                "payload": {
                    "name": "Engram is a conditional memory module designed to augment Transformer backbones.",
                    "section_path": "Engram",
                },
                "evidence": evidence,
            },
            {
                "local_id": "concept-engram",
                "object_type": "concept",
                "payload": {"name": "Engram", "section_path": "Engram"},
                "evidence": evidence,
            },
            {
                "local_id": "claim-ground",
                "object_type": "claim",
                "payload": {
                    "name": "The quiet ground return path avoids noise coupling.",
                    "section_path": "Engram",
                },
                "evidence": evidence,
            },
        ]
        relations = [
            {
                "source_local_id": "claim-engram",
                "target_local_id": "concept-engram",
                "edge_type": "about",
                "evidence": [{"quote": evidence[0]["quoted_span"]}],
            }
        ]
        n_obj, n_rel = repo.store_kg(nb.id, source_id, objects, relations)
        assert (n_obj, n_rel) == (3, 1)

        graph = repo.knowledge_graph(nb.id)
        assert len(graph.nodes) == 3 and len(graph.edges) == 1
        types = {item.object_type: item.count for item in repo.knowledge_types(nb.id)}
        assert types.get("claim") == 2 and types.get("concept") == 1
        claims = repo.list_knowledge(nb.id, "claim").items
        assert claims and claims[0].headline

        answer = repo.ask(
            nb.id,
            AskRequest(question="What is Engram conditional memory module?", mode="reasoning"),
        )
        assert answer.answer_id
        assert answer.conversation_id
        assert answer.llm_mode == "deterministic"
        assert answer.related_knowledge
        assert answer.citations

        feedback = repo.submit_feedback(
            answer.answer_id,
            FeedbackRequest(rating="useful", comment="grounded"),
        )
        assert feedback.answer_id == answer.answer_id

        conversations = repo.list_conversations(nb.id)
        assert len(conversations) == 1 and conversations[0].turn_count == 1
        detail = repo.get_conversation(answer.conversation_id)
        assert len(detail.turns) == 1
        repo.rename_conversation(answer.conversation_id, "Engram thread")
        assert repo.get_conversation(answer.conversation_id).title == "Engram thread"
        second = repo.ask(
            nb.id,
            AskRequest(
                question="Does the Engram conditional memory module separate storage and computation?",
                conversation_id=answer.conversation_id,
                mode="reasoning",
            ),
        )
        assert second.conversation_id == answer.conversation_id
        assert repo.get_conversation(answer.conversation_id).turn_count == 2
        repo.delete_conversation(answer.conversation_id)
        assert repo.list_conversations(nb.id) == []

        claim_id = next(item.id for item in repo.list_knowledge(nb.id, "claim").items)
        repo.update_knowledge(nb.id, claim_id, KnowledgeUpdate(status="deprecated", owner="curator"))
        assert all(node.id != claim_id for node in repo.knowledge_graph(nb.id).nodes)
        dep_answer = repo.ask(
            nb.id,
            AskRequest(question="What is Engram conditional memory module?", mode="reasoning"),
        )
        assert all(item.id != claim_id for item in dep_answer.related_knowledge)
        repo.update_knowledge(nb.id, claim_id, KnowledgeUpdate(status="reviewed"))
        reviewed = next(item for item in repo.list_knowledge(nb.id, "claim").items if item.id == claim_id)
        assert reviewed.status == "reviewed" and reviewed.last_reviewed


def check_api_layer() -> None:
    """Boot the FastAPI app and exercise route groups against a throwaway DB."""
    from fastapi.testclient import TestClient

    with tempfile.TemporaryDirectory(prefix="sn-api-") as tmp:
        os.environ.update(
            {
                "MODEL_SERVICES_CONFIG": "",
                "MINERU_MODE": "off",
                "DATABASE_URL": f"sqlite:///{tmp}/api.db",
                "SILICON_NOTEBOOK_STORAGE_DIR": f"{tmp}/storage",
                "LLM_LOG_ENABLED": "false",
                "EVENT_LOG_ENABLED": "false",
                "SILICON_NOTEBOOK_AUTH_OPTIONAL": "true",
            }
        )
        from app.api.deps import repository as route_repository
        from app.core import readiness
        from app.core.config import get_settings
        from app.main import create_app

        get_settings.cache_clear()
        route_repository.cache_clear()
        # This route smoke intentionally drives TestClient without its lifespan,
        # just like the pytest API suite. Startup/readiness warm-up has dedicated
        # tests; mark this isolated throwaway app ready before exercising routes.
        readiness.mark_ready()
        client = TestClient(create_app())

        def ok(method: str, path: str, **kw):
            response = client.request(method, path, **kw)
            assert response.status_code == 200, (
                f"{method} {path} -> {response.status_code}: {response.text[:200]}"
            )
            return response.json()

        ok("GET", "/api/health")
        ok("GET", "/api/me")
        doc_types = ok("GET", "/api/doc-types")
        assert {item["id"] for item in doc_types} >= {"", "academic_paper", "textbook"}

        nb = ok(
            "POST",
            "/api/notebooks",
            json={
                "name": "API smoke",
                "purpose": "p",
                "primary_domain": "d",
            },
        )
        nid = nb["id"]
        assert ok("GET", "/api/notebooks")
        ok("GET", f"/api/notebooks/{nid}")
        _srcs = ok("GET", f"/api/notebooks/{nid}/sources")  # 分页 shape:{items,total_count,...}
        assert _srcs["items"] == [] and _srcs["total_count"] == 0

        # Two-tier: set tier to base, then back to personal, via the REST route.
        assert ok("GET", f"/api/notebooks/{nid}")["tier"] == "personal"
        set_base = ok("POST", f"/api/notebooks/{nid}/tier", json={"tier": "base"})
        assert set_base["tier"] == "base"
        assert ok("GET", f"/api/notebooks/{nid}")["tier"] == "base"
        set_personal = ok("POST", f"/api/notebooks/{nid}/tier", json={"tier": "personal"})
        assert set_personal["tier"] == "personal"
        # Bad tier -> 400; missing notebook -> 404.
        assert client.post(f"/api/notebooks/{nid}/tier", json={"tier": "bogus"}).status_code == 400
        assert client.post("/api/notebooks/nb-missing/tier", json={"tier": "base"}).status_code == 404

        ok("GET", f"/api/notebooks/{nid}/knowledge-types")
        ok("GET", f"/api/notebooks/{nid}/knowledge", params={"type": "claim"})
        ok("GET", f"/api/notebooks/{nid}/duplicates", params={"type": "claim"})
        ok("GET", f"/api/notebooks/{nid}/graph")
        ok("GET", f"/api/notebooks/{nid}/analytics")

        schemas = ok("GET", "/api/object-schemas")
        assert any(schema["object_type"] == "concept" for schema in schemas)
        ok(
            "POST",
            "/api/object-schemas",
            json={"object_type": "api_test_type", "fields": ["name"], "label": "x"},
        )
        ok("PATCH", "/api/object-schemas/api_test_type", json={"status": "disabled"})
        assert client.delete("/api/object-schemas/api_test_type").status_code == 200
        ok("POST", f"/api/notebooks/{nid}/schema-proposals")

        # PR#334:空库(无来源/无库内知识/无挂载参考库)对 /ask 会被后端以 409 硬拒
        # (ask_available=False 的权威闸门)。本 smoke 只为走通 API 层的问答链路,先经
        # 客户端共享的路由仓库单例、走**公有** upload_sources 传一条最小来源(离线可
        # 同步完成解析/切块,与 check_kg_layer 同法),使该笔记本有可检索来源 →
        # ask_available=True。刻意不直插行/不碰私有 _write seam,免踩仓库边界守卫。
        from app.services.repository import UploadedSourceFile

        route_repository().upload_sources(
            nid,
            [
                UploadedSourceFile(
                    file_name="seed.md",
                    content_type="text/markdown",
                    content=(
                        "# Seed\n\nMinimal retrievable evidence so the API-layer "
                        "ask clears the empty-notebook gate.\n"
                    ).encode("utf-8"),
                )
            ],
        )

        answer = ok(
            "POST",
            f"/api/notebooks/{nid}/ask",
            json={"question": "any notebook knowledge?", "scenario": {}},
        )
        assert "related_knowledge" in answer and answer["llm_mode"] == "deterministic"
        assert answer["conversation_id"]
        conversations = ok("GET", f"/api/notebooks/{nid}/conversations")
        assert len(conversations) == 1
        conversation_id = answer["conversation_id"]
        detail = ok("GET", f"/api/conversations/{conversation_id}")
        assert len(detail["turns"]) == 1
        ok("PATCH", f"/api/conversations/{conversation_id}", json={"title": "renamed"})
        ok("DELETE", f"/api/conversations/{conversation_id}")

        assert client.get("/api/notebooks/nb-does-not-exist").status_code == 404
        assert client.get("/api/notebooks/nb-does-not-exist/knowledge", params={"type": "claim"}).status_code == 404
        assert client.patch("/api/object-schemas/no_such_type", json={"label": "x"}).status_code == 404
        assert client.delete("/api/object-schemas/concept").status_code == 400
        dup = client.post("/api/object-schemas", json={"object_type": "concept", "fields": ["name"]})
        assert dup.status_code == 400
        assert client.get(f"/api/notebooks/{nid}/knowledge").status_code == 422

        route_repository.cache_clear()
        get_settings.cache_clear()


def check_csv_xlsx_parsing() -> None:
    """CSV and XLSX parse into table_row elements."""
    from app.services.parsers import parse_csv, parse_xlsx

    with tempfile.TemporaryDirectory(prefix="sn-csv-") as tmp:
        csv_path = Path(tmp) / "rules.csv"
        csv_path.write_text(
            "Pin,Net,Note\n1,GND,quiet return\n2,VIN,sensitive input\n",
            encoding="utf-8",
        )
        csv_elements = parse_csv("src-csv", csv_path)
        assert len(csv_elements) == 3
        assert all(element.element_type == "table_row" for element in csv_elements)
        assert "GND" in csv_elements[1].text

        try:
            from openpyxl import Workbook
        except ImportError:
            return
        xlsx_path = Path(tmp) / "rules.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Pin", "Net", "Note"])
        sheet.append([1, "GND", "quiet return"])
        workbook.save(str(xlsx_path))
        xlsx_elements = parse_xlsx("src-xlsx", xlsx_path)
        assert xlsx_elements and xlsx_elements[0].element_type == "table_row"
        assert "GND" in " ".join(element.text for element in xlsx_elements)


def check_mineru_mapping() -> None:
    """MinerU content_list -> SourceElements: formulas keep LaTeX, tables flatten."""
    content_list = [
        {"type": "title", "text": "Bondwire coupling", "text_level": 1, "page_idx": 0},
        {"type": "text", "text": "Return path matters.", "text_level": 0, "page_idx": 0},
        {
            "type": "equation",
            "text": "$$\nV = I \\cdot R\n$$",
            "text_format": "latex",
            "page_idx": 1,
        },
        {
            "type": "table",
            "table_body": "<table><tr><td>Pin</td><td>Net</td></tr><tr><td>1</td><td>GND</td></tr></table>",
            "table_caption": ["Pinout"],
            "page_idx": 1,
        },
    ]
    elements = mineru_content_list_to_elements("src-x", content_list)
    types = {element.element_type for element in elements}
    assert "heading" in types and "paragraph" in types
    formula = next(element for element in elements if element.element_type == "formula")
    assert "V = I \\cdot R" in formula.text and "$$" not in formula.text
    assert formula.location_label == "PDF p.2 formula 1"
    table = next(element for element in elements if element.element_type == "table")
    assert "Pinout" in table.text and "GND" in table.text
    assert table.metadata.get("table_html", "").startswith("<table>")


def check_score_knowledge() -> None:
    """score_knowledge: keyword baseline unchanged, vector adds semantic recall."""
    objects = [
        {
            "id": "k1",
            "payload": {"name": "quiet ground return path"},
            "evidence": [_ev("e1", "quiet ground")],
        },
        {
            "id": "k2",
            "payload": {"name": "thermal via array"},
            "evidence": [_ev("e2", "thermal via")],
        },
    ]

    keyword_only = score_knowledge("quiet ground", objects, "claim")
    assert keyword_only and keyword_only[0].object_id == "k1"

    element_vectors = {"e1": [1.0, 0.0], "e2": [0.0, 1.0]}
    semantic = score_knowledge("xyzzy plugh", objects, "claim", [0.0, 1.0], element_vectors)
    assert semantic and semantic[0].object_id == "k2"

    none_vector = score_knowledge("quiet ground", objects, "claim", None, element_vectors)
    assert [item.object_id for item in none_vector] == [item.object_id for item in keyword_only]


def check_cjk_keyword_recall() -> None:
    """Chinese keyword retrieval works without verbatim phrase equality."""
    from app.services.retrieval import _tokens, keyword_score

    tokens = _tokens("低噪声封装")
    assert {"低噪", "噪声", "声封", "封装"} <= set(tokens), tokens
    assert keyword_score("低噪声 封装", "本节讨论低噪声模拟前端的封装注意事项") > 0

    claims = [
        {
            "id": "ko-cjk",
            "payload": {"name": "低噪声模拟前端封装应保证安静地回流路径。"},
            "evidence": [],
            "status": "approved",
        },
        {
            "id": "ko-other",
            "payload": {"name": "高速数字时钟走线需要包地处理。"},
            "evidence": [],
            "status": "approved",
        },
    ]
    scored = score_knowledge("低噪声 封装 回流", claims, "claim", None, None)
    ids = [item.object_id for item in scored]
    assert ids and ids[0] == "ko-cjk", ids
    assert "ko-other" not in ids
    print("[smoke] cjk keyword recall ok")


def check_hybrid_floor_and_weight() -> None:
    """Relevance floor drops noise; type weight stays separate from relevance."""
    objects = [
        {"id": "hit", "payload": {"name": "quiet ground return path"}, "evidence": [], "status": "approved"},
        {"id": "noise", "payload": {"name": "thermal vias for power module layout"}, "evidence": [], "status": "approved"},
    ]
    scored = score_knowledge("quiet ground return", objects, "claim", None, None)
    ids = [item.object_id for item in scored]
    assert "hit" in ids and "noise" not in ids, ids
    assert scored[0].weight == 1.0
    assert abs(scored[0].score - scored[0].relevance) < 1e-9
    print("[smoke] hybrid floor + weight ok")


def check_payload_embedding_retrieval() -> None:
    """Payload-level vectors can beat unrelated evidence vectors."""
    query_vector = [1.0, 0.0]
    knowledge_vectors = {"A": [1.0, 0.0], "B": [0.0, 1.0]}
    element_vectors = {"elB": [0.0, 1.0]}
    objects = [
        {"id": "A", "payload": {"name": "alpha"}, "evidence": [], "status": "approved"},
        {
            "id": "B",
            "payload": {"name": "beta"},
            "evidence": [type("E", (), {"element_id": "elB", "quoted_span": "z"})()],
            "status": "approved",
        },
    ]
    scored = score_knowledge(
        "zzz",
        objects,
        "claim",
        query_vector,
        element_vectors,
        knowledge_vectors,
    )
    assert scored and scored[0].object_id == "A", [item.object_id for item in scored]
    print("[smoke] payload-level embedding retrieval ok")


def check_json_fences() -> None:
    """chat_json strips ```json fences some models wrap responses in."""
    from app.core.llm import strip_json_fences

    assert strip_json_fences('```json\n{"a": 1}\n```') == '{"a": 1}'
    assert strip_json_fences('```\n{"b": 2}\n```') == '{"b": 2}'
    assert strip_json_fences('{"c": 3}') == '{"c": 3}'


def check_llm_interaction_logging() -> None:
    """LLM interaction logger writes parseable JSONL and never propagates IO errors."""
    from app.core.event_logging import get_log_owner, owner_dir
    from app.core.llm_logging import LLMInteractionLogger

    with tempfile.TemporaryDirectory(prefix="silicon-notebook-llmlog-") as temp_dir:
        log_path = Path(temp_dir) / "logs" / "llm.jsonl"
        logger = LLMInteractionLogger(Settings(llm_log_path=str(log_path), llm_log_max_chars=20))
        clipped = logger.clip("x" * 100)
        assert clipped.startswith("x" * 20) and "chars]" in clipped, clipped
        logger.log(
            {
                "ts": "2026-05-29T00:00:00",
                "id": "llm-test",
                "kind": "chat",
                "model": "test-model",
                "status": "ok",
                "latency_ms": 12,
                "usage": {"total_tokens": 7},
                "response": {"content": "{}"},
            }
        )
        # per-user 模式(PR#93):llm 日志写在 dirname(llm_log_path)/<owner>/ 下,
        # owner 未设(离线/无请求上下文)→ user-local。
        # 注意:EventLogger 现在按天分文件 <channel>-YYYY-MM-DD.jsonl(配套 gzip 归档),
        # 已不再是单个 llm.jsonl——logger.filename/.path 只是历史兼容属性,emit() 实际写的是
        # 当天的 llm-<date>.jsonl。故这里按天文件名(channel 恒为 "llm")glob 读实际落盘那份,
        # 别再假设旧的 llm.jsonl 命名(那正是本 check 之前 FileNotFoundError 的根因)。
        owner_log_dir = log_path.parent / owner_dir(get_log_owner())
        dated = sorted(owner_log_dir.glob("llm-*.jsonl"))
        assert dated, f"per-user llm 日志未落盘于 {owner_log_dir}"
        written = dated[-1]
        line = written.read_text(encoding="utf-8").strip().splitlines()[-1]
        parsed = json.loads(line)
        for field in ("ts", "id", "kind", "model", "status", "latency_ms"):
            assert field in parsed, field

        # 韧性:emit 走 _target_path_for_day(不读 .path),故旧的 `bad.path=.../nope.jsonl`
        # 覆盖已失效、测不到吞错。改用「日志目录的父路径是文件」真正触发写失败:emit 必须
        # 吞掉 IO 错、绝不上抛。
        clash = Path(temp_dir) / "clash"
        clash.write_text("x", encoding="utf-8")
        bad = LLMInteractionLogger(Settings(llm_log_path=str(clash / "llm.jsonl")))
        bad.log({"kind": "embed", "model": "m", "status": "ok", "latency_ms": 1, "dims": 3})

        # 禁用:enabled=False → 完全不写(per-user 模式连目录都不建),故目录不应存在。
        # (旧断言 `not off.jsonl.exists()` 恒真——enabled 时也只会写 off-<date>.jsonl、
        #  从不写 off.jsonl,测不到禁用;改断言实际写入目录不存在。)
        off_dir = Path(temp_dir) / "off-logs"
        off = LLMInteractionLogger(
            Settings(llm_log_path=str(off_dir / "llm.jsonl"), llm_log_enabled=False))
        off.log({"kind": "chat", "model": "m", "status": "ok", "latency_ms": 1})
        assert not off_dir.exists()
    print("[smoke] llm interaction logging ok")


def check_event_logging() -> None:
    """EventLogger writes parseable JSONL, honors disable, and never raises."""
    from app.core.event_logging import EventLogger

    with tempfile.TemporaryDirectory(prefix="silicon-notebook-eventlog-") as temp_dir:
        logger = EventLogger(Settings(event_log_dir=temp_dir), channel="events")
        logger.emit({"kind": "pipeline", "stage": "parse", "status": "done", "latency_ms": 5})
        # EventLogger 按天分文件 events-YYYY-MM-DD.jsonl(配套 gzip 归档),已不再是单个
        # events.jsonl;glob 读实际落盘那份。
        dated = sorted(Path(temp_dir).glob("events-*.jsonl"))
        assert dated, f"events 日志未落盘于 {temp_dir}"
        line = dated[-1].read_text(encoding="utf-8").strip().splitlines()[-1]
        parsed = json.loads(line)
        assert parsed["channel"] == "events" and "ts" in parsed, parsed
        assert parsed["stage"] == "parse"

        # 韧性:log_dir 指向一个文件时建目录/写入必失败,但 emit 必须吞掉、绝不上抛。
        # (旧的 `bad.path=.../nope.jsonl` 覆盖已失效——emit 不读 .path。)
        clash = Path(temp_dir) / "clash-ev"
        clash.write_text("x", encoding="utf-8")
        bad = EventLogger(Settings(event_log_dir=str(clash)), channel="events")
        bad.emit({"kind": "status", "status": "ok"})

        # 禁用:enabled=False → 无 off-<date>.jsonl 落盘(旧的 not off.jsonl.exists() 恒真)。
        off = EventLogger(Settings(event_log_dir=temp_dir, event_log_enabled=False), channel="off")
        off.emit({"kind": "x", "status": "ok"})
        assert not list(Path(temp_dir).glob("off*.jsonl"))
    print("[smoke] event logging ok")


def check_pipeline_event_logging() -> None:
    """process_source emits stage events and stores real failure error_message."""
    from app.core.event_logging import get_log_owner, owner_dir

    with tempfile.TemporaryDirectory(prefix="silicon-notebook-pipe-") as temp_dir:
        root = Path(temp_dir)
        log_dir = root / "logs"
        repo = SQLiteRepository(_offline_settings(root, event_log_dir=log_dir))
        notebook = repo.create_notebook(NotebookCreate(name="Pipe", purpose="p", primary_domain="d"))
        uploaded = repo.upload_sources(
            notebook.id,
            [
                UploadedSourceFile(
                    file_name="s.md",
                    content_type="text/markdown",
                    content=b"# H\n\nrequire a quiet ground return path.\n",
                )
            ],
        )
        assert uploaded[0].parse_status == "extracted"
        run = _latest_extraction_run(repo, uploaded[0].id)
        assert run["run_type"] == "kg" and run["error_message"] == "no-llm"

        # per-user 模式(PR#93):仓库 event_log 写在 event_log_dir/<owner>/ 下;离线无请求
        # 上下文 → owner=user-local。EventLogger 按天分文件 events-YYYY-MM-DD.jsonl,glob 读取
        # 实际落盘的那些(单次运行通常一份;跨天则合并读)。
        owner_ev_dir = log_dir / owner_dir(get_log_owner())
        events_files = sorted(owner_ev_dir.glob("events-*.jsonl"))
        assert events_files, f"pipeline events 未落盘于 {owner_ev_dir}"
        events = [
            json.loads(line)
            for path in events_files
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        pipeline = [event for event in events if event.get("kind") == "pipeline"]
        stages = {event["stage"] for event in pipeline}
        assert {"parse", "embed", "extract"} <= stages, stages
        assert any(event["stage"] == "pipeline" and event["status"] == "done" for event in pipeline)
        assert any(event.get("kind") == "status" for event in events)

        source = repo.get_source(uploaded[0].id)
        Path(source.file_path).unlink()
        result = repo.parse_source(uploaded[0].id)
        assert result.parse_status == "failed"
        assert result.error_message
        assert result.error_message != "Parsing failed; see source error."
    print("[smoke] pipeline event logging + error_message ok")


def main() -> None:
    check_csv_xlsx_parsing()
    check_mineru_mapping()
    check_score_knowledge()
    check_extraction_profiles()
    check_object_schemas()
    check_kg_record_binding()
    check_kg_extract_window_grounding()
    check_kg_windowing()
    check_kg_store_ask_and_conversations()
    check_json_fences()
    check_llm_interaction_logging()
    check_event_logging()
    check_pipeline_event_logging()
    check_cjk_keyword_recall()
    check_hybrid_floor_and_weight()
    check_payload_embedding_retrieval()
    check_api_layer()

    with tempfile.TemporaryDirectory(prefix="silicon-notebook-smoke-") as temp_dir:
        root = Path(temp_dir)
        repository = SQLiteRepository(_offline_settings(root))

        assert repository.list_notebooks() == []
        assert repository.current_user().display_name

        notebook = repository.create_notebook(
            NotebookCreate(
                name="Smoke notebook",
                purpose="Validate local persistence and parsing.",
                primary_domain="Semiconductor",
            )
        )

        auto_nb = repository.create_notebook(NotebookCreate(name="Auto desc"))
        assert auto_nb.purpose == ""
        manual_nb = repository.create_notebook(NotebookCreate(name="Manual", purpose="hand-written"))
        assert manual_nb.purpose == "hand-written"

        # Two-tier: tier toggle round-trips on the repository.
        assert repository.get_notebook(manual_nb.id).tier == "personal"
        repository.mark_notebook_base(manual_nb.id)
        assert repository.get_notebook(manual_nb.id).tier == "base"
        repository.set_notebook_personal(manual_nb.id)
        assert repository.get_notebook(manual_nb.id).tier == "personal"
        try:
            repository.set_notebook_personal("nb-missing")
            raise AssertionError("set_notebook_personal should raise KeyError on missing notebook")
        except KeyError:
            pass
        typed = repository.upload_sources(
            auto_nb.id,
            [
                UploadedSourceFile(
                    file_name="paper.md",
                    content_type="text/markdown",
                    content=b"# Abstract\nWe propose X and report experimental results. References: [1].\n",
                    doc_type="academic_paper",
                )
            ],
        )
        assert typed[0].doc_type == "academic_paper"
        assert _latest_extraction_run(repository, typed[0].id)["error_message"] == "no-llm"
        # L4: 导入后从来源自动生成描述(purpose_auto=1 时)；smoke 无 LLM，走确定性兜底。
        auto_purpose = repository.get_notebook(auto_nb.id).purpose
        assert auto_purpose and "本笔记本收录了" in auto_purpose
        edited = repository.update_notebook(auto_nb.id, NotebookUpdate(purpose="pinned"))
        assert edited.purpose == "pinned"
        repository.upload_sources(
            auto_nb.id,
            [
                UploadedSourceFile(
                    file_name="more.md",
                    content_type="text/markdown",
                    content=b"# More\nExtra content here.\n",
                )
            ],
        )
        assert repository.get_notebook(auto_nb.id).purpose == "pinned"

        uploaded = repository.upload_sources(
            notebook.id,
            [
                UploadedSourceFile(
                    file_name="mixed_markdown.md",
                    content_type="text/markdown",
                    content=(
                        "# Wirebond review\n\n"
                        "Engram is a conditional memory module for Transformer backbones.\n\n"
                        "低噪声模拟前端需要检查 quiet ground return path.\n\n"
                        "- ESD path must not share sensitive input return.\n"
                    ).encode("utf-8"),
                ),
                UploadedSourceFile(
                    file_name="review_notes.docx",
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    content=_docx_bytes("DOCX parser smoke paragraph for package review."),
                ),
                UploadedSourceFile(
                    file_name="slides.pptx",
                    content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    content=_pptx_bytes("PPTX parser smoke slide for pin assignment."),
                ),
                UploadedSourceFile(
                    file_name="empty.pdf",
                    content_type="application/pdf",
                    content=_pdf_bytes(),
                ),
            ],
        )

        assert len(uploaded) == 4
        assert all(source.parse_status == "extracted" for source in uploaded)
        assert all(_latest_extraction_run(repository, source.id)["run_type"] == "kg" for source in uploaded)

        markdown = next(source for source in uploaded if source.file_name.endswith(".md"))
        assert len(repository.source_elements(markdown.id)) >= 3

        docx = next(source for source in uploaded if source.file_name.endswith(".docx"))
        assert repository.source_elements(docx.id)[0].text.startswith("DOCX parser smoke")

        pptx = next(source for source in uploaded if source.file_name.endswith(".pptx"))
        pptx_elements = repository.source_elements(pptx.id)
        assert "PPTX parser smoke" in pptx_elements[0].text
        assert any(element.element_type == "slide_text" and "Second shape" in element.text for element in pptx_elements)
        assert any(element.element_type == "speaker_notes" for element in pptx_elements)

        hits = repository.search_notebook(notebook.id, "quiet ground")
        assert any(hit.scope == "Element" for hit in hits.hits)

        evidence = [_source_evidence(repository, markdown.id)]
        kg_objects = [
            {
                "local_id": "claim-1",
                "object_type": "claim",
                "payload": {
                    "name": "Engram is a conditional memory module for Transformer backbones.",
                    "section_path": "Wirebond review",
                },
                "evidence": evidence,
            },
            {
                "local_id": "concept-1",
                "object_type": "concept",
                "payload": {"name": "Engram", "section_path": "Wirebond review"},
                "evidence": evidence,
            },
        ]
        kg_relations = [
            {
                "source_local_id": "claim-1",
                "target_local_id": "concept-1",
                "edge_type": "about",
                "evidence": [{"quote": evidence[0]["quoted_span"]}],
            }
        ]
        assert repository.store_kg(notebook.id, markdown.id, kg_objects, kg_relations) == (2, 1)

        ktypes = {item.object_type: item.count for item in repository.knowledge_types(notebook.id)}
        assert ktypes.get("claim") == 1 and ktypes.get("concept") == 1
        generic_claims = repository.list_knowledge(notebook.id, "claim").items
        assert generic_claims and any(field.key == "name" for field in generic_claims[0].fields)

        answer = repository.ask(
            notebook.id,
            AskRequest(question="What is Engram conditional memory module?", mode="reasoning"),
        )
        assert answer.answer_id
        assert answer.related_knowledge
        assert answer.citations
        assert answer.llm_mode == "deterministic"

        feedback = repository.submit_feedback(
            answer.answer_id,
            FeedbackRequest(rating="useful", comment="grounded and actionable"),
        )
        assert feedback.answer_id == answer.answer_id
        assert feedback.comment == "grounded and actionable"

        claim_id = generic_claims[0].id
        repository.update_knowledge(notebook.id, claim_id, KnowledgeUpdate(status="deprecated", owner="curator-a"))
        dep_answer = repository.ask(
            notebook.id,
            AskRequest(question="What is Engram conditional memory module?", mode="reasoning"),
        )
        assert all(item.id != claim_id for item in dep_answer.related_knowledge)
        repository.update_knowledge(notebook.id, claim_id, KnowledgeUpdate(status="reviewed"))
        reviewed = next(item for item in repository.list_knowledge(notebook.id, "claim").items if item.id == claim_id)
        assert reviewed.status == "reviewed" and reviewed.last_reviewed
        try:
            repository.update_knowledge(notebook.id, claim_id, KnowledgeUpdate(status="bogus"))
            raise AssertionError("invalid status should raise ValueError")
        except ValueError:
            pass

        rule_id = _insert_rule(
            repository,
            notebook.id,
            markdown.id,
            title="Keep quiet ground return path",
            statement="Low-noise analog inputs should keep a quiet ground return path.",
            applies_to=["wirebond", "analog input"],
            recommendation="Route ESD return separately from sensitive input return.",
            risk_if_ignored="Noise can couple into the analog front-end.",
        )
        rule_id_2 = _insert_rule(
            repository,
            notebook.id,
            markdown.id,
            title="Keep quiet return path",
            statement="Analog inputs should preserve a quiet ground return path.",
            applies_to=["wirebond", "analog input"],
            recommendation="Route the return separately.",
        )
        rules_now = repository.list_knowledge(notebook.id, "rule").items
        assert {rule.id for rule in rules_now} >= {rule_id, rule_id_2}
        assert isinstance(repository.list_knowledge(notebook.id, "method").items, list)
        assert isinstance(repository.list_knowledge(notebook.id, "risk").items, list)
        assert isinstance(repository.list_knowledge(notebook.id, "glossary").items, list)
        assert isinstance(repository.find_duplicates(notebook.id, "rule"), list)

        merged = repository.merge_knowledge(notebook.id, rule_id_2, MergeRequest(into_id=rule_id))
        assert merged.id == rule_id
        assert any(
            rule.id == rule_id_2 and rule.status == "deprecated"
            for rule in repository.list_knowledge(notebook.id, "rule").items
        )

        analytics = repository.notebook_analytics(notebook.id)
        assert analytics.feedback_useful >= 1
        assert analytics.usefulness_rate == 1.0
        assert analytics.knowledge_counts.get("rule", 0) >= 1
        assert sum(analytics.source_status_counts.values()) >= 1

        reparsed = repository.parse_source(markdown.id)
        assert reparsed.parse_status == "extracted"
        assert all(item.id != claim_id for item in repository.list_knowledge(notebook.id, "claim").items)
        assert all(rule.id != rule_id for rule in repository.list_knowledge(notebook.id, "rule").items)

        scheduled: list[str] = []
        queued = repository.upload_sources(
            notebook.id,
            [
                UploadedSourceFile(
                    file_name="deferred.md",
                    content_type="text/markdown",
                    content=b"# Deferred\n\nrequire post-upload background parsing.\n",
                )
            ],
            scheduler=scheduled.append,
        )
        assert queued[0].parse_status == "queued"
        assert repository.source_elements(queued[0].id) == []
        assert scheduled == [queued[0].id]
        processed = repository.process_source(scheduled[0])
        assert processed.parse_status == "extracted"
        assert repository.source_elements(queued[0].id)
        processed_path = Path(repository.get_source(processed.id).file_path)
        assert processed_path.exists()
        repository.delete_source(processed.id)
        assert all(source.id != processed.id for source in repository.list_sources(notebook.id))
        assert not processed_path.exists()
        try:
            repository.source_elements(processed.id)
            raise AssertionError("deleted source should no longer expose elements")
        except KeyError:
            pass

        restarted = SQLiteRepository(_offline_settings(root))
        assert restarted.get_notebook(notebook.id).name == "Smoke notebook"
        assert restarted.source_elements(markdown.id)

    print("backend smoke ok: sqlite persistence, KG extraction boundary, retrieval, feedback")


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    slide_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p>"
        "</p:txBody></p:sp>"
        "<p:sp><p:txBody><a:p><a:r><a:t>Second shape risk note</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    notes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r>"
        "<a:t>Speaker note: confirm pin map before signoff.</a:t>"
        "</a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:notes>"
    )
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", slide_xml)
        archive.writestr("ppt/notesSlides/notesSlide1.xml", notes_xml)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    from pypdf import PdfWriter

    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    main()
